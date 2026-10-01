"""Regressions reproduced from the eight-document Agent comparison."""
import json
from types import SimpleNamespace
from unittest.mock import AsyncMock

import fitz
import pytest
from PIL import Image
from test_agentic_repair import Model, config, schedule

from app.pipeline.agentic import AgenticRepair, RegionModel, Verification, visual_issues
from app.pipeline.agentic_content import replacement_text
from app.pipeline.agentic_evidence import EvidenceSource, match_word_pages
from app.pipeline.final_delivery import finalize_delivery, preserve_document_content
from app.pipeline.stages.chunk import ChunkStage
from app.pipeline.stages.package import PackageStage


class InspectModel(Model):
    def __init__(self, *, keep=False, omit_last=False, grouped=False):
        super().__init__()
        self.keep = keep
        self.omit_last = omit_last
        self.grouped = grouped
        self.contexts = []

    async def query(self, task, payload, image, schema):
        self.contexts.append((task, payload, image))
        result, tokens = await super().query(task, payload, image, schema)
        if task == "decide" and self.keep:
            result["action"] = "keep"
        if task == "repair":
            if self.omit_last:
                result["facts"][0]["source_cells"] = [
                    c for c in payload["required_source_cells"] if not c.startswith("r3")
                ]
            if self.grouped:
                result["facts"].append({**result["facts"][0], "object": "第二個可核對的事實"})
        return result, tokens


@pytest.mark.asyncio
async def test_decision_sees_both_renderings_and_verified_keep_avoids_rewrite(tmp_path):
    model = InspectModel(keep=True)
    ir = schedule(tmp_path)
    ir.blocks[0].payload["table_body"] = ir.blocks[0].payload["table_body"].replace(
        "<td>工作</td>", '<td colspan="2">工作</td>'
    ).replace("<td>市場調查</td><td></td>", "<td>市場調查</td><td>1月至3月</td>")
    output, report = await AgenticRepair(config(), model, model).run(ir, tmp_path)
    assert report["kept"] == 1 and report["accepted"] == 0
    assert model.calls == ["repair", "decide", "verify"]
    context = next(payload for task, payload, _ in model.contexts if task == "decide")
    assert "<table>" in context["source_html"]
    assert "第 2 列：第 1 欄：市場調查；第 2 欄：1月至3月" in context["current_markdown"]
    assert context["current_rag_markdown"]
    assert output.blocks[0].payload["agentic_needs_review"] is False
    assert not visual_issues(output)
    assert not replacement_text(output.blocks[0])
    await AgenticRepair(config(), model, model).run(ir, tmp_path)
    assert len(model.calls) == 3


@pytest.mark.asyncio
async def test_one_observation_can_explicitly_verify_multiple_facts(tmp_path):
    model = InspectModel(grouped=True)
    output, report = await AgenticRepair(config(), model, model).run(schedule(tmp_path), tmp_path)
    assert report["accepted"] == 1
    assert len(output.blocks[0].payload["agentic_evidence"]["facts"]) == 2


@pytest.mark.asyncio
async def test_repeated_task_label_cannot_hide_missing_last_row(tmp_path):
    ir = schedule(tmp_path)
    ir.blocks[0].payload["table_body"] = ir.blocks[0].payload["table_body"].replace(
        "</table>", "<tr><td>市場調查</td>" + "<td></td>" * 6 + "</tr></table>"
    )
    model = InspectModel(omit_last=True)
    output, report = await AgenticRepair(config(max_attempts_per_block=1), model, model).run(ir, tmp_path)
    assert report["accepted"] == 0
    assert model.calls == ["repair"]
    assert output.blocks[0].payload["agentic_needs_review"]
    state = json.loads((tmp_path / "agentic/state.json").read_text())
    assert "r3c1" in state["events"][-1]["error"]


def test_verifier_must_map_each_fact_to_its_own_cells():
    from app.pipeline.agentic import Interpretation

    candidate = Interpretation(kind="data_table", needs_review=False, facts=[
        {"subject": "A", "relation": "值", "object": "1", "source_cells": ["r1c1"]},
        {"subject": "B", "relation": "值", "object": "2", "source_cells": ["r2c1"]},
    ])
    review = Verification(complete=True, needs_review=False, missing_source_cells=[], observations=[
        {"fact_indices": [1], "source_cells": ["r1c1"], "observation": "A值1"},
        {"fact_indices": [0], "source_cells": ["r2c1"], "observation": "B值2"},
    ])
    with pytest.raises(ValueError, match="mapping"):
        AgenticRepair._check_review(review, candidate, {"r1c1", "r2c1"}, {"r1c1", "r2c1"})


@pytest.mark.asyncio
async def test_cross_page_evidence_and_hash_invalidation(tmp_path):
    ir = schedule(tmp_path)
    next_page = tmp_path / "assets/pages/p0001.png"
    Image.new("RGB", (600, 500), "blue").save(next_page)
    model = InspectModel()
    await AgenticRepair(config(), model, model).run(ir, tmp_path)
    assert len(model.contexts[0][2]) == 2
    assert model.contexts[0][2][0].parent.name == "agentic"
    assert model.contexts[0][2][1] == next_page
    Image.new("RGB", (600, 500), "red").save(next_page)
    await AgenticRepair(config(), model, model).run(ir, tmp_path)
    assert len(model.calls) == 6


@pytest.mark.asyncio
async def test_complete_repair_has_one_delivery_and_keeps_original_evidence(tmp_path):
    ir = schedule(tmp_path)
    ir.blocks[0].payload["table_body"] = ir.blocks[0].payload["table_body"].replace("市場調查", "錯誤配對原文")
    model = Model()
    cfg = config()
    cfg.enrich.enable_vlm = False
    cfg.package.enable_semantic_repair = False
    repaired, report = await AgenticRepair(cfg, model, model).run(ir, tmp_path)
    assert report["accepted"] == 1
    result = await PackageStage(cfg).run("doc", "run", repaired, tmp_path)
    assert result.success, result.error
    result = await ChunkStage(cfg).run("doc", "run", repaired, tmp_path)
    assert result.success, result.error
    main = (tmp_path / "outputs/documents/main.md").read_text()
    chunks = (tmp_path / "outputs/chunks.jsonl").read_text()
    assert main.count("市場調查｜執行月份：1月至3月") == 1
    assert chunks.count("市場調查｜執行月份：1月至3月") == 1
    assert "錯誤配對原文" not in main and "錯誤配對原文" not in chunks
    assert "錯誤配對原文" in (tmp_path / "outputs/dataset.md").read_text()
    assert "錯誤配對原文" in repaired.blocks[0].payload["table_body"]
    assert preserve_document_content(repaired, tmp_path) == 0
    status = finalize_delivery(repaired, tmp_path, cfg)
    assert status["coverage"]["missing_documents"] == 0
    assert status["coverage"]["missing_chunks"] == 0


@pytest.mark.asyncio
async def test_truncated_response_counts_tokens_and_informs_retry(tmp_path):
    from app.config import VLMApiMode

    response = SimpleNamespace(
        choices=[SimpleNamespace(message=object(), finish_reason="length")],
        usage=SimpleNamespace(total_tokens=1900),
    )
    adapter = SimpleNamespace(
        config=SimpleNamespace(model="test", api_mode=VLMApiMode.OLLAMA),
        client=SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=AsyncMock(return_value=response)))),
        _message_text=lambda _: '{"action":',
        _extract_json_payload=lambda v: v,
        _encode_image_as_data_uri=lambda _: "data:image/png;base64,",
    )
    model = RegionModel(adapter, 4096)
    ir = schedule(tmp_path)
    _, report = await AgenticRepair(config(max_model_calls=2), model, model).run(ir, tmp_path)
    assert report["usage"]["tokens"] == 3800
    state = json.loads((tmp_path / "agentic/state.json").read_text())
    assert all(c["finish_reason"] == "length" for c in state["calls"])
    second_prompt = adapter.client.chat.completions.create.call_args_list[1].kwargs["messages"][1]["content"][0]["text"]
    assert "output_truncated" in second_prompt


def test_word_table_alignment_requires_distinctive_contiguous_source(tmp_path):
    ir = schedule(tmp_path)
    table = ir.blocks[0]
    table.payload["table_body"] = "<table><tr><td>品牌永續策略</td></tr><tr><td>顧客訪談研究</td></tr></table>"
    assert match_word_pages(table, ["其他內容", "品牌永續策略", "顧客訪談研究"]) == [1, 2]
    assert match_word_pages(table, ["品牌永續策略顧客訪談研究", "目錄頁", "無關頁",
                                    "品牌永續策略顧客訪談研究"]) == []


@pytest.mark.asyncio
async def test_word_pdf_evidence_preserves_logical_ir_and_embedded_image_paths(tmp_path, monkeypatch):
    from app.pipeline.stages.parse import ParseStage
    ir = schedule(tmp_path)
    ir.page_kind = "logical"
    ir.source.ext = "docx"
    ir.blocks[0].payload["table_body"] = "<table><tr><td>First unique label</td></tr><tr><td>Second unique label</td></tr></table>"
    source = tmp_path / "original.docx"
    source.write_bytes(b"mock docx input")

    async def convert(source, output):
        output.mkdir(parents=True, exist_ok=True)
        pdf = output / "original.pdf"
        with fitz.open() as doc:
            page = doc.new_page()
            page.insert_text((20, 20), "First unique label\nSecond unique label")
            doc.save(pdf)
        return pdf, None

    monkeypatch.setattr(ParseStage, "_convert_office_to_pdf", staticmethod(convert))
    evidence = EvidenceSource(ir, tmp_path)
    await evidence.prepare_word(source)
    files = evidence.for_block(ir.blocks[0])
    assert files and files[0][0] == "Word converted PDF page 1"
    assert ir.page_kind == "logical" and ir.blocks[0].page_idx == 0
    cache = tmp_path / "cache"
    (cache / "native/images").mkdir(parents=True)
    Image.new("RGB", (50, 50)).save(cache / "native/images/embedded.png")
    block = ir.blocks[0].model_copy(deep=True)
    block.type = "image"
    block.payload = {"img_path": "images/embedded.png"}
    evidence = EvidenceSource(ir, tmp_path, cache)
    assert evidence.for_block(block)[0][0] == "Original Word embedded image"
    block.payload["img_path"] = "../native/images/embedded.png"
    assert not evidence.for_block(block)


@pytest.mark.asyncio
async def test_schedule_headers_are_retained_without_inventing_header_facts(tmp_path):
    class HeaderModel(Model):
        async def query(self, task, payload, image, schema):
            value, tokens = await super().query(task, payload, image, schema)
            if task == "repair":
                value["facts"][0]["source_cells"] = payload["required_fact_cells"]
                assert "current_markdown" not in payload and "source_html" not in payload
            return value, tokens

    model = HeaderModel()
    output, report = await AgenticRepair(config(), model, model).run(schedule(tmp_path), tmp_path)
    assert report["accepted"] == 1 and report["usage"]["model_calls"] == 3
    assert "欄位：工作｜1月｜2月" in replacement_text(output.blocks[0])
    assert "市場調查" in replacement_text(output.blocks[0])


@pytest.mark.asyncio
async def test_reviewer_findings_reach_next_decision_and_repair(tmp_path):
    class FeedbackModel(Model):
        def __init__(self):
            super().__init__()
            self.reviews = 0

        async def query(self, task, payload, image, schema):
            value, tokens = await super().query(task, payload, image, schema)
            if task == "verify":
                self.reviews += 1
                if self.reviews == 1:
                    value["complete"] = False
                    value["observations"][0]["observation"] = "月份關係對錯列"
            if task == "repair" and self.reviews:
                assert "月份關係對錯列" in json.dumps(payload["previous_attempts"], ensure_ascii=False)
                assert "current_markdown" not in payload
            return value, tokens

    model = FeedbackModel()
    _, report = await AgenticRepair(config(), model, model).run(schedule(tmp_path), tmp_path)
    assert report["accepted"] == 1 and model.reviews == 2


def test_cell_reference_cannot_hide_omitted_task_details():
    from app.pipeline.agentic import Interpretation
    candidate = Interpretation(kind="schedule", needs_review=False, facts=[{
        "subject": "項目1內外部調查", "relation": "執行月份", "object": "6月", "source_cells": ["r2c1"],
    }])
    with pytest.raises(ValueError, match="source_text_missing"):
        AgenticRepair._check_candidate(candidate, {"r2c1"}, {"r2c1"}, [{
            "id": "r2c1", "text": "項目1內外部調查 1.企業內部需求訪談 2.市場分析",
        }])


def test_word_alignment_weights_distinctive_header_over_repeated_work_names(tmp_path):
    ir = schedule(tmp_path)
    ir.blocks[0].payload["table_body"] = "<table>" + "".join(
        f"<tr><td>{label}</td></tr>" for label in
        ["特定時程表頭", "品牌永續策略", "永續溝通策略", "永續溝通應用", "品牌體驗設計", "市場分析研究", "顧客需求調查"]
    ) + "</table>"
    repeated = "品牌永續策略永續溝通策略永續溝通應用品牌體驗設計市場分析研究顧客需求調查"
    assert match_word_pages(ir.blocks[0], [repeated, "", "", repeated, "", "", "特定時程表頭" + repeated]) == [6]


def test_invalid_legacy_png_is_not_sent_as_raster_evidence(tmp_path):
    ir = schedule(tmp_path)
    ir.page_kind = "logical"
    block = ir.blocks[0]
    block.type = "image"
    block.payload = {"img_path": "images/vector.png"}
    cache = tmp_path / "cache"
    (cache / "images").mkdir(parents=True)
    (cache / "images/vector.png").write_bytes(b"not a raster PNG")
    assert EvidenceSource(ir, tmp_path, cache).for_block(block) == []


def test_isolated_word_picture_retains_the_exact_requested_image(tmp_path):
    import hashlib

    from docx import Document

    from app.pipeline.agentic_evidence import _isolate_word_picture

    first, second = tmp_path / "first.png", tmp_path / "second.png"
    Image.new("RGB", (20, 20), "red").save(first)
    Image.new("RGB", (20, 20), "blue").save(second)
    document = Document()
    document.add_paragraph("Unrelated text")
    document.add_picture(str(first))
    document.add_picture(str(second))
    source, fragment = tmp_path / "source.docx", tmp_path / "picture.docx"
    document.save(source)
    assert _isolate_word_picture(source, hashlib.sha256(second.read_bytes()).hexdigest(), fragment)
    isolated = Document(fragment)
    assert len(isolated.inline_shapes) == 1
    assert not any(p.text for p in isolated.paragraphs)
    shape = isolated.inline_shapes[0]
    part = isolated.part.related_parts[shape._inline.graphic.graphicData.pic.blipFill.blip.embed]
    assert part.blob == second.read_bytes()


def test_corrected_column_can_relocate_text_within_the_same_source_row():
    from app.pipeline.agentic import Interpretation
    candidate = Interpretation(kind="data_table", needs_review=False, facts=[
        {"subject": "健康意識", "relation": "說明", "object": "台灣45%上班族長期痠痛、超過62%自評身體亞健康", "source_cells": ["r3c2"]},
        {"subject": "健康意識", "relation": "意義", "object": "需求已經從醫療預防日常保養", "source_cells": ["r3c3"]},
    ])
    AgenticRepair._check_candidate(candidate, {"r3c2", "r3c3"}, {"r3c2", "r3c3"}, [
        {"id": "r3c2", "row": 3, "text": "超過62%自評身體亞健康"},
        {"id": "r3c3", "row": 3, "text": "台灣45%上班族長期痰痛、需求已經從醫療預防日常保養"},
    ])


def test_multiline_headers_are_retained_and_cannot_hide_late_data_rows(tmp_path):
    from app.pipeline.agentic import Interpretation
    from app.pipeline.agentic_evidence import table_cells

    ir = schedule(tmp_path)
    block = ir.blocks[0]
    block.payload["table_body"] = "<table><tr><td colspan='2'>標題</td></tr><tr><td>項目</td><td>說明</td></tr><tr><td>工作</td><td>資料內容</td></tr><tr><td>第二項</td><td>資料內容二</td></tr></table>"
    cells = table_cells(block)
    candidate = Interpretation(kind="data_table", needs_review=False, header_cells=["r1c1", "r2c1", "r2c2"], facts=[
        {"subject": "工作", "relation": "說明", "object": "資料內容", "source_cells": ["r3c1", "r3c2"]},
    ])
    assert len(AgenticRepair._retained_headers(candidate, cells, [])) == 3
    candidate.header_cells = ["r4c1", "r4c2"]
    with pytest.raises(ValueError, match="leading three"):
        AgenticRepair._retained_headers(candidate, cells, [])


@pytest.mark.asyncio
async def test_figure_facts_have_explicit_valid_image_references(tmp_path):
    class FigureModel(Model):
        async def query(self, task, payload, image, schema):
            value, tokens = await super().query(task, payload, image, schema)
            if task == "decide":
                value["action"] = "describe_figure"
            if task == "repair":
                assert payload["allowed_source_ids"] == ["image:1"]
                assert payload["image_sources"][0]["id"] == "image:1"
            return value, tokens

    ir = schedule(tmp_path)
    from app.models.document_ir import BlockType
    ir.blocks[0].type = BlockType.IMAGE
    ir.blocks[0].payload = {"caption": "Figure"}
    model = FigureModel()
    output, report = await AgenticRepair(config(), model, model).run(ir, tmp_path)
    assert report["accepted"] == 1
    assert output.blocks[0].payload["agentic_evidence"]["facts"][0]["source_cells"] == ["image:1"]


def test_keep_cannot_override_a_verified_column_correction():
    from app.pipeline.agentic import Interpretation
    candidate = Interpretation(kind="data_table", needs_review=False, facts=[{
        "subject": "健康意識提升", "relation": "說明",
        "object": "台灣45%上班族長期痠痛、超過62%自評身體亞健康", "source_cells": ["r3c1", "r3c2"],
    }])
    assert not AgenticRepair._existing_cells_match(candidate, [
        {"id": "r3c1", "text": "健康意識提升"}, {"id": "r3c2", "text": "超過62%自評身體亞健康"},
        {"id": "r3c3", "text": "台灣45%上班族長期痰痛、需求已經從醫療預防日常保養"},
    ])


@pytest.mark.asyncio
async def test_explicit_page_inspection_keeps_crop_and_correct_page_labels(tmp_path):
    class ContextModel(InspectModel):
        async def query(self, task, payload, image, schema):
            value, tokens = await super().query(task, payload, image, schema)
            decisions = sum(t == "decide" for t, _, _ in self.contexts)
            if task == "decide" and decisions == 1:
                value["action"] = "inspect_page"
            if task == "repair" and decisions:
                assert len(image) == 3
                assert payload["evidence_images"] == [
                    "PDF page 1 (target crop)", "PDF page 1 (full context)", "PDF page 2",
                ]
            return value, tokens

    ir = schedule(tmp_path)
    Image.new("RGB", (600, 500), "blue").save(tmp_path / "assets/pages/p0001.png")
    model = ContextModel()
    _, report = await AgenticRepair(config(), model, model).run(ir, tmp_path)
    assert report["accepted"] == 1 and report["usage"]["model_calls"] == 5


def test_shared_header_verification_applies_to_each_data_fact():
    from app.pipeline.agentic import Interpretation
    candidate = Interpretation(kind="data_table", needs_review=False, facts=[
        {"subject": "營收", "relation": "2025年", "object": "100", "source_cells": ["r1c1", "r2c1"]},
        {"subject": "成本", "relation": "2025年", "object": "80", "source_cells": ["r1c1", "r3c1"]},
    ])
    review = Verification(complete=True, needs_review=False, missing_source_cells=[], observations=[
        {"fact_indices": [], "source_cells": ["r1c1"], "observation": "共同欄位表頭是2025年"},
        {"fact_indices": [0], "source_cells": ["r2c1"], "observation": "營收列數值100"},
        {"fact_indices": [1], "source_cells": ["r3c1"], "observation": "成本列數值80"},
    ])
    refs = {"r1c1", "r2c1", "r3c1"}
    AgenticRepair._check_review(review, candidate, refs, refs, shared_headers={"r1c1"})
    with pytest.raises(ValueError, match="mapping"):
        AgenticRepair._check_review(review, candidate, refs, refs)
