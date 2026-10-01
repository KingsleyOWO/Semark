"""Integration: a real native Word pipeline, persistent queue intent, and resume."""
import asyncio
import hashlib
import json

import pytest
from docx import Document
from PIL import Image

from app.config import PROFILES, ProfileName, settings
from app.core.orchestrator import PipelineOrchestrator, reset_pipeline_semaphore
from app.core.submissions import has_submission, save_submission
from app.db.database import Database
from app.db.repositories import DocRepository, RunRepository, RunStageRepository
from app.models.entities import DocCreate, RunCreate, RunStatus, StageName, StageStatus
from app.pipeline.stages.enrich import enrichment_input_hash
from app.pipeline.stages.normalize import load_document_ir


async def seed(db, tmp_path):
    source = settings.get_doc_path("doc") / "source/original.docx"
    source.parent.mkdir(parents=True)
    document = Document()
    document.add_heading("年度計畫成果", level=1)
    document.add_paragraph("本計畫以提升產品品質與顧客滿意度為目標，執行期間完成市場調查與服務改善。")
    document.add_paragraph("數位行銷預算 1,100,000 元。")
    document.save(source)
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    await DocRepository(db).create("doc", DocCreate(
        source_path=str(source), sha256=digest, ext="docx", size_bytes=source.stat().st_size))
    config = PROFILES[ProfileName.FAST].model_copy(deep=True)
    config.package.enable_semantic_repair = False
    run = await RunRepository(db).create(RunCreate(
        doc_id="doc", profile="fast", config=config.model_dump(mode="json"), config_hash="config"))
    await RunStageRepository(db).create_all_stages(run.run_id)
    return run


@pytest.mark.asyncio
async def test_real_pipeline_publishes_complete_revision_and_reuses_checkpoint(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "workspace_path", tmp_path / "workspace")
    reset_pipeline_semaphore()
    db = Database(tmp_path / "test.db")
    await db.connect()
    try:
        run = await seed(db, tmp_path)
        result = await PipelineOrchestrator(db).execute(run.run_id)
        assert result.success, result.error
        root = settings.get_run_path("doc", run.run_id)
        published = settings.get_published_run_path("doc", run.run_id)
        assert (root / "current").is_symlink()
        assert published.parent == root / "revisions"
        ir = load_document_ir(published)
        assert ir.page_kind == "logical"
        status = json.loads((published / "outputs/delivery_status.json").read_text())
        assert status["coverage"]["missing_documents"] == 0
        assert status["coverage"]["missing_chunks"] == 0
        assert "1,100,000" in (root / "outputs/documents/main.md").read_text()
        assert (published / "document_ir.baseline.json").exists()

        async def must_not_parse(*args, **kwargs):
            raise AssertionError("Completed parse must not repeat")

        before = await RunStageRepository(db).list_by_run(run.run_id)
        resumed = PipelineOrchestrator(db)
        monkeypatch.setattr(resumed, "_run_parse", must_not_parse)
        result = await resumed.execute(run.run_id)
        assert result.success, result.error
        assert settings.get_published_run_path("doc", run.run_id) == published
        after = await RunStageRepository(db).list_by_run(run.run_id)
        assert [(s.started_at, s.finished_at) for s in after] == [
            (s.started_at, s.finished_at) for s in before]
        assert all(s.stats["resumed_from_checkpoint"] for s in after if s.stage != StageName.INGEST)
    finally:
        await db.disconnect()
        reset_pipeline_semaphore()


@pytest.mark.asyncio
async def test_failure_after_enrich_resumes_at_package(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "workspace_path", tmp_path / "workspace")
    reset_pipeline_semaphore()
    db = Database(tmp_path / "test.db")
    await db.connect()
    try:
        run = await seed(db, tmp_path)
        interrupted = PipelineOrchestrator(db)

        async def fail_package(*args, **kwargs):
            raise RuntimeError("simulated package interruption")

        monkeypatch.setattr(interrupted, "_run_package", fail_package)
        result = await interrupted.execute(run.run_id)
        assert not result.success
        root = settings.get_run_path("doc", run.run_id)
        state = json.loads((root / "checkpoint.json").read_text())
        assert state["completed"] == ["parse", "normalize", "enrich"]
        assert not (root / "current").exists()

        async def must_not_enrich(*args, **kwargs):
            raise AssertionError("Completed enrichment must not repeat")

        resumed = PipelineOrchestrator(db)
        real_package = resumed._run_package

        async def check_retry_timers(*args, **kwargs):
            stage = await RunStageRepository(db).get(run.run_id, StageName.PACKAGE)
            assert stage.status == StageStatus.RUNNING
            assert stage.started_at is not None
            assert stage.finished_at is None
            assert stage.error is None
            return await real_package(*args, **kwargs)

        monkeypatch.setattr(resumed, "_run_enrich", must_not_enrich)
        monkeypatch.setattr(resumed, "_run_package", check_retry_timers)
        result = await resumed.execute(run.run_id)
        assert result.success, result.error
        assert (root / "current").exists()
    finally:
        await db.disconnect()
        reset_pipeline_semaphore()


@pytest.mark.asyncio
async def test_startup_resumes_submitted_work_only(tmp_path, monkeypatch):
    import app.main as main

    monkeypatch.setattr(settings, "workspace_path", tmp_path / "workspace")
    db = Database(tmp_path / "test.db")
    await db.connect()
    monkeypatch.setattr(main, "db", db)
    try:
        run = await seed(db, tmp_path)
        save_submission(run)
        assert has_submission(run)
        assert await main._cancel_orphan_runs() == 0

        class Queue:
            submitted = []
            async def submit(self, run_id):
                self.submitted.append(run_id)

        queue = Queue()
        assert await main._resume_submitted_runs(queue) == 1
        assert queue.submitted == [run.run_id]
        assert (await RunRepository(db).get(run.run_id)).status == RunStatus.PENDING
    finally:
        await db.disconnect()


def test_enrichment_input_hash_changes_for_actual_text_and_image(tmp_path):
    from app.models.document_ir import Block, BlockType, DocumentIR, EngineInfo, SourceInfo
    ir = DocumentIR(doc_id="doc", run_id="run",
        source=SourceInfo(path="a.pdf", ext="pdf", sha256="source", size_bytes=1),
        engine=EngineInfo(backend="pipeline", method="auto"),
        blocks=[Block(block_id="b", type=BlockType.TEXT, page_idx=0, payload={"text": "預算100元"})])
    first = enrichment_input_hash(ir, tmp_path, None)
    other_run = ir.model_copy(update={"run_id": "another"})
    assert enrichment_input_hash(other_run, tmp_path, None) == first
    other_run = ir.model_copy(deep=True)
    other_run.blocks[0].payload["text"] = "預算200元"
    assert enrichment_input_hash(other_run, tmp_path, None) != first
    image = tmp_path / "assets/pages/p0000.png"
    image.parent.mkdir(parents=True)
    Image.new("RGB", (60, 60), "white").save(image)
    before = enrichment_input_hash(ir, tmp_path, None)
    Image.new("RGB", (60, 60), "black").save(image)
    assert enrichment_input_hash(ir, tmp_path, None) != before



@pytest.mark.asyncio
async def test_concurrent_execution_of_same_run_only_parses_once(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "workspace_path", tmp_path / "workspace")
    reset_pipeline_semaphore()
    db = Database(tmp_path / "test.db")
    await db.connect()
    calls = []
    original = PipelineOrchestrator._run_parse

    async def counted(self, *args, **kwargs):
        calls.append(1)
        await asyncio.sleep(0.02)
        return await original(self, *args, **kwargs)

    monkeypatch.setattr(PipelineOrchestrator, "_run_parse", counted)
    try:
        run = await seed(db, tmp_path)
        results = await asyncio.gather(
            PipelineOrchestrator(db).execute(run.run_id),
            PipelineOrchestrator(db).execute(run.run_id),
        )
        assert all(result.success for result in results)
        assert len(calls) == 1
        root = settings.get_run_path("doc", run.run_id)
        assert len(list((root / "revisions").iterdir())) == 1
    finally:
        await db.disconnect()
        reset_pipeline_semaphore()


@pytest.mark.asyncio
async def test_missing_parse_cache_rebuilds_without_mutating_published_revision(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "workspace_path", tmp_path / "workspace")
    reset_pipeline_semaphore()
    db = Database(tmp_path / "test.db")
    await db.connect()
    try:
        run = await seed(db, tmp_path)
        result = await PipelineOrchestrator(db).execute(run.run_id)
        assert result.success, result.error
        root = settings.get_run_path("doc", run.run_id)
        first = settings.get_published_run_path("doc", run.run_id)
        baseline = {str(p.relative_to(first)): p.read_bytes() for p in first.rglob("*") if p.is_file()}
        checkpoint = json.loads((root / "checkpoint.json").read_text())
        from pathlib import Path
        Path(checkpoint["content_list_path"]).unlink()
        result = await PipelineOrchestrator(db).execute(run.run_id)
        assert result.success, result.error
        assert settings.get_published_run_path("doc", run.run_id) != first
        assert all((first / name).read_bytes() == content for name, content in baseline.items())
    finally:
        await db.disconnect()
        reset_pipeline_semaphore()
