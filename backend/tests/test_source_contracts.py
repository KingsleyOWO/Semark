"""Source relationships survive assembly; context and protocol errors cannot certify content."""
import json

import pytest
from docx import Document
from PIL import Image
from test_document_recovery import RecoveryModel, setup
from test_region_recovery import table_source

from app.adapters.word import parse_word_document
from app.models.document_ir import Block, BlockType
from app.pipeline.agentic_evidence import table_cells
from app.pipeline.document_review import PageReview
from app.pipeline.image_reading import read_image_parts_bounded, read_native_batch
from app.pipeline.package_utils import semantic_table_to_text
from app.pipeline.recovery_models import RecoveryDecision, RegionReading, RegionText
from app.pipeline.region_recovery import locate_prose, validate_region_rows
from app.pipeline.stages.chunk import ChunkStage
from app.pipeline.stages.enrich import EnrichStage
from app.pipeline.stages.normalize import NormalizeStage, save_document_ir
from app.pipeline.stages.package import PackageStage
from app.pipeline.table_structure import (
    TableCell,
    TableStructure,
    attach_structure,
    table_structure,
)


def test_grid_keeps_merged_coordinates_empty_cells_and_distinct_repeated_values():
    b = Block(block_id='t', type=BlockType.TABLE, page_idx=2, bbox_norm=[0, 10, 1000, 500], payload={
        'table_body': '<table><tr><th rowspan="2">Metric</th><th colspan="2">Year</th></tr>'
                      '<tr><th>2025</th><th>2026</th></tr><tr><td>Revenue</td><td>100</td><td></td></tr>'
                      '<tr><td>Revenue</td><td>100</td><td>120</td></tr></table>'})
    before = b.model_dump()
    grid = table_structure(b)
    assert b.model_dump() == before  # Read-only evidence collection.
    assert next(c for c in grid.cells if c.id == 'r2c2').text == '2025'
    assert next(c for c in grid.cells if c.id == 'r3c3').text == ''
    assert len([c for c in grid.cells if c.text == 'Revenue']) == 2
    attach_structure(b, grid)
    assert 'rowspan="2"' in b.get_text() and 'colspan="2"' in b.get_text()
    assert next(c for c in table_cells(b) if c['id'] == 'r2c3')['text'] == '2026'
    b.payload['table_body'] = '<table><tr><td>New</td><td>240</td></tr></table>'
    assert table_structure(b).cells[1].text == '240'  # A stale stored grid cannot mask a changed body.


def test_recovered_table_does_not_promote_first_data_record_to_header():
    grid = TableStructure(table_id='work', row_count=3, column_count=3, cells=[
        TableCell(row=1, column=1, text='輔導架構內容', header=True),
        TableCell(row=1, column=2, text='', header=True),
        TableCell(row=1, column=3, text='預期完成產出', header=True),
        TableCell(row=2, column=1, text='A【品牌定位建構】'),
        TableCell(row=2, column=2, text='內外部資源盤點'), TableCell(row=2, column=3, text='研究報告'),
        TableCell(row=3, column=1, text='B【品牌網站】'),
        TableCell(row=3, column=2, text='網站1式'), TableCell(row=3, column=3, text='無後台')])
    result = semantic_table_to_text(grid.to_html(), '完成產出')
    assert '輔導架構內容：B【品牌網站】' in result
    assert 'A【品牌定位建構】：B【品牌網站】' not in result


class TableModel(RecoveryModel):
    async def query(self, task, payload, image, schema):
        if task == 'read_region':
            return {'text': '', 'rows': [
                {'source_id': r['source_id'], 'role': 'header' if r['row'] == 1 else 'data',
                 'cells': [{'column': c['column'], 'text': next(
                     (n['text'] for n in r['native_cells'] if n['column'] == c['column']), '')}
                           for c in r['cells']]}
                for r in payload['owned_rows']], 'complete': True, 'observation': 'Source rows preserved.'}, 20
        return await super().query(task, payload, image, schema)


@pytest.mark.asyncio
async def test_regions_assemble_one_table_with_final_row_and_source_coordinates(tmp_path):
    ir, owner, runner = await setup(tmp_path, TableModel())
    inventory = table_source(tmp_path)
    ir.blocks = inventory.blocks()
    runner.inventory.pdf = inventory.pdf
    runner.inventory.ir = ir
    runner.inventory.mapping = inventory.mapping
    result, _ = await runner.read_regions(ir, {'page_idx': 0},
        RecoveryDecision(action='read_regions', reason='Read source rows'), tmp_path / 'candidate', '')
    assert len(result.blocks) == 1 and result.blocks[0].type == BlockType.TABLE
    grid = table_structure(result.blocks[0])
    assert [c.text for c in grid.cells if c.column == 1] == ['Metric', 'Revenue', 'Profit', 'Growth']
    assert [c.text for c in grid.cells if c.column == 2] == ['Value', '100', '80', '20']
    assert all(c.source_bbox for c in grid.cells)
    assert grid.source['source_rows'][-1] == 'native:r4'
    rendered = ChunkStage(owner.config)._block_to_text(result.blocks[0])
    assert '- Metric：Growth' in rendered and '- Value：20' in rendered
    save_document_ir(result, tmp_path / 'saved')
    saved = json.loads((tmp_path / 'saved/document_ir.json').read_text())
    assert saved['blocks'][0]['payload']['table_structure']['source']['source_rows'][-1] == 'native:r4'


@pytest.mark.parametrize('kind', ['missing', 'duplicate', 'unowned', 'wrong_column', 'wrong_value', 'free_markdown'])
def test_region_contract_rejects_loss_duplicate_and_misattribution(kind):
    source = {'source_id': 'table:r2', 'row': 2, 'cells': [{'column': 1}, {'column': 2}],
              'native_cells': [{'id': 'r2c1', 'column': 1, 'text': 'Revenue', 'header': False},
                               {'id': 'r2c2', 'column': 2, 'text': '100', 'header': False}]}
    row = {'source_id': 'table:r2', 'role': 'data', 'cells': [{'column': 1, 'text': 'Revenue'}, {'column': 2, 'text': '100'}]}
    data = {'rows': [row], 'complete': True, 'observation': 'Claimed complete'}
    if kind == 'missing':
        data.update(rows=[], text='omitted table')
    elif kind == 'duplicate':
        data['rows'] = [row, row]
    elif kind == 'unowned':
        row['source_id'] = 'table:r1'
    elif kind == 'wrong_column':
        row['cells'][1]['column'] = 1
    elif kind == 'wrong_value':
        row['cells'][1]['text'] = '200'
    else:
        data['text'] = '| Revenue | 100 |'
    with pytest.raises(ValueError, match='region_'):
        validate_region_rows({'owned_rows': [source]}, RegionReading.model_validate(data))


def test_word_paragraph_caption_is_context_and_cannot_skip_image_reading(tmp_path):
    path = tmp_path / 'picture.png'
    Image.new('RGB', (80, 80), 'green').save(path)
    doc = Document()
    neighbor = 'Figure: this nearby paragraph describes a different subject.'
    doc.add_paragraph(neighbor)
    doc.add_picture(str(path))
    source = tmp_path / 'source.docx'
    doc.save(source)
    parsed = parse_word_document(source, tmp_path / 'parsed')
    items = json.loads(parsed.content_list_path.read_text())
    block = next(NormalizeStage()._parse_block(item, i) for i, item in enumerate(items) if item['type'] == 'image')
    assert block.payload['caption_context'] == neighbor
    assert block.get_text() == ''
    assert block.payload['word_source']['caption_origin'] == 'neighboring_text'
    assert EnrichStage(None)._get_enrichment_kind(block, None, False) == 'figure_description'
    assert neighbor not in ChunkStage(None)._block_to_text(block)
    legacy = block.model_copy(deep=True)
    legacy.payload.pop('caption_context')
    legacy.payload['caption'] = neighbor
    assert legacy.get_text() == ''  # Old IRs use the same provenance rule.


@pytest.mark.asyncio
async def test_planner_sees_actual_tool_limits_before_spending_repair_actions(tmp_path):
    class Observe(RecoveryModel):
        async def query(self, task, payload, image, schema):
            if task == 'recovery_plan':
                self.payload = payload
                return {'action': 'defer', 'reason': 'Source ownership prevents a whole page edit.'}, 10
            return await super().query(task, payload, image, schema)
    model = Observe()
    ir, owner, runner = await setup(tmp_path, model)
    runner.inventory.mapping['b0'] = [0, 1]
    result = await runner.run(ir)
    assert result is ir and runner.state['actions'] == 0
    for tool in ['read_regions', 'read_source', 'reparse_pdf', 'rebuild_delivery']:
        assert tool not in model.payload['tools']
        assert not model.payload['tool_constraints'][tool]['available']
    assert model.payload['source_features']['block_count'] == 1
    with pytest.raises(ValueError, match='tool_precondition_failed'):
        await runner.read_regions(ir, {'page_idx': 0}, RecoveryDecision(action='read_regions', reason='Invalid'), tmp_path, '')
    assert model.tasks == []


@pytest.mark.asyncio
async def test_native_format_failure_splits_group_and_reuses_successful_reads(tmp_path):
    class Split(RecoveryModel):
        def __init__(self):
            super().__init__()
            self.groups = []
        async def query(self, task, payload, image, schema):
            self.groups.append(payload['image_order'])
            if len(payload['sources']) > 1:
                return {'items': []}, 10
            sid = payload['sources'][0]['source_id']
            return {'items': [{'source_id': sid, 'text': 'Text for ' + sid, 'complete': True,
                               'observation': 'Read the source.'}]}, 10
    model = Split()
    _, owner, runner = await setup(tmp_path, model)
    picture = runner.inventory.image(0)
    payload = {'sources': [{'source_id': 'a'}, {'source_id': 'b'}], 'image_order': ['a', 'b']}
    for _ in range(2):
        result = await read_native_batch(owner, 'verify_native_reading', payload, [picture, picture], source_ids=['a', 'b'])
        assert [i.text for i in result.items] == ['Text for a', 'Text for b']
    assert model.groups == [['a', 'b'], ['a'], ['b']]


@pytest.mark.asyncio
async def test_image_region_overflow_uses_two_crops_and_keeps_original_coordinates(tmp_path):
    class Parts(RecoveryModel):
        def __init__(self):
            super().__init__()
            self.scopes = []
        async def query(self, task, payload, image, schema):
            core = payload.get('source_core')
            self.scopes.append(core)
            if not core:
                return {'image_type': 'document', 'visual_facts': [], 'visual_independent_of_text': False,
                        'text_regions': [], 'complete': True, 'observation': 'Content omitted by invalid response'}, 10
            return {'image_type': 'document', 'visual_facts': [], 'visual_independent_of_text': False,
                    'text_regions': [{'region_id': 'line', 'bbox': [0, 0, 1000, 1000], 'readable': True,
                                      'text': 'Top' if core[1] == 0 else 'Bottom'}],
                    'complete': True, 'observation': 'Source crop content'}, 10
    model = Parts()
    _, owner, runner = await setup(tmp_path, model)
    picture = runner.inventory.image(0)
    for _ in range(2):
        result = await read_image_parts_bounded(owner, 'read_image_parts', {}, picture)
        assert [r.text for r in result.text_regions] == ['Top', 'Bottom']
        assert result.text_regions[0].bbox[1] == 0 and result.text_regions[1].bbox[3] == 1000
        assert result.text_regions[0].bbox[3] == result.text_regions[1].bbox[1]
    assert len(model.scopes) == 3


@pytest.mark.asyncio
async def test_audit_format_split_keeps_whole_page_issues_and_does_not_certify_missing_content(tmp_path):
    class Audit(RecoveryModel):
        async def query(self, task, payload, image, schema):
            self.tasks.append(task)
            ids = payload['required_block_ids']
            result = {'blocks': [{'block_id': bid, 'status': 'verified', 'context_complete': True,
                                  'observation': 'x' * 701 if len(ids) == 2 else 'Source agrees'} for bid in ids],
                      'source_complete': False, 'reading_order_correct': True,
                      'observation': 'A source paragraph has no block.', 'issues': ['Missing source paragraph']}
            return result, 10
    model = Audit()
    _, owner, runner = await setup(tmp_path, model)
    payload = {'required_block_ids': ['a', 'b'], 'blocks': [{'block_id': 'a'}, {'block_id': 'b'}],
               'delivered_documents': [{'text': 'Actual output'}], 'delivered_chunks': []}
    result = await owner._audit_page_contract(payload, [runner.inventory.image(0)])
    assert isinstance(result, PageReview) and len(result.blocks) == 2
    assert not result.source_complete and result.issues == ['Missing source paragraph']
    assert model.tasks == ['audit_page'] * 3


def test_short_table_values_survive_final_main_document_cleaning():
    grid = TableStructure(table_id='schedule', row_count=3, column_count=3, cells=[
        TableCell(row=1, column=1, text='工作', header=True),
        TableCell(row=1, column=2, text='6月', header=True),
        TableCell(row=1, column=3, text='11月', header=True),
        TableCell(row=2, column=1, text='調查'), TableCell(row=2, column=2, text='●'),
        TableCell(row=2, column=3, text=''), TableCell(row=3, column=1, text='結案'),
        TableCell(row=3, column=2, text=''), TableCell(row=3, column=3, text='●')])
    semantic = semantic_table_to_text(grid.to_html(), '時程')
    main = PackageStage()._clean_split_main_body(semantic, '時程', [])
    assert '- 6月：●' in semantic and '- 6月：●' in main
    assert '- 11月：●' in semantic and '- 11月：●' in main
    assert not PackageStage._is_low_value_visual_ocr_line('- A: 0')
    assert PackageStage._is_low_value_visual_ocr_line('x_')


@pytest.mark.asyncio
async def test_audit_cache_is_invalidated_when_actual_delivery_changes(tmp_path):
    model = RecoveryModel()
    _, owner, runner = await setup(tmp_path, model)
    payload = {'required_block_ids': ['b0'], 'blocks': [{'block_id': 'b0'}],
               'delivered_documents': [{'text': 'Revenue 100'}], 'delivered_chunks': []}
    images = [runner.inventory.image(0)]
    await owner._audit_page_contract(payload, images)
    await owner._audit_page_contract(payload, images)
    assert model.tasks.count('audit_page') == 1
    payload['delivered_documents'][0]['text'] = 'Revenue 200'
    await owner._audit_page_contract(payload, images)
    assert model.tasks.count('audit_page') == 2


def test_prose_location_uses_native_anchor_and_rejects_table_fragments():
    prose = RegionText(text='Heading', bbox=[10, 300, 250, 340])
    native = [('heading', [10, 80, 250, 100]), ('revenue100', [20, 320, 250, 340])]
    cells = [[10, 200, 500, 400]]
    bbox, origin = locate_prose(prose, [0, 0, 1000, 500], native, cells)
    assert bbox == [10, 80, 250, 100] and origin == 'native_text'
    duplicate = RegionText(text='Revenue 100', bbox=[10, 60, 250, 100])
    with pytest.raises(ValueError, match='region_prose_overlaps_table_cell'):
        locate_prose(duplicate, [0, 0, 1000, 500], native, cells)


@pytest.mark.asyncio
async def test_text_repair_tools_cannot_flatten_structured_table(tmp_path):
    class Flatten(RecoveryModel):
        async def query(self, task, payload, image, schema):
            return {'complete': True, 'edits': [{'block_id': 'b0', 'text': 'Revenue 100',
                    'evidence_ids': payload['source_evidence_ids'], 'reason': 'Flatten source table'}]}, 10
    ir, _, runner = await setup(tmp_path, Flatten())
    ir.blocks[0].type = BlockType.TABLE
    ir.blocks[0].payload = {'table_body': '<table><tr><td>Revenue</td><td>100</td></tr></table>'}
    constraints = runner.tool_constraints({'page_idx': 0}, ir.blocks)
    assert not constraints['read_source']['available']
    assert constraints['rebuild_delivery']['text_edit_block_ids'] == []
    with pytest.raises(ValueError, match='structured blocks cannot become free text'):
        await runner.rebuild_delivery(ir, {'page_idx': 0},
            RecoveryDecision(action='rebuild_delivery', reason='Fix'), tmp_path, '')
    assert ir.blocks[0].type == BlockType.TABLE


@pytest.mark.asyncio
async def test_region_format_failure_splits_at_rows_and_preserves_complete_table(tmp_path):
    class SplitTable(TableModel):
        def __init__(self):
            super().__init__()
            self.read_rows = []
        async def query(self, task, payload, image, schema):
            if task == 'read_region':
                ids = [r['source_id'] for r in payload['owned_rows']]
                self.read_rows.append(ids)
                if len(ids) > 1:
                    return {'complete': True, 'observation': 'Truncated response with no rows'}, 10
            return await super().query(task, payload, image, schema)
    model = SplitTable()
    ir, _, runner = await setup(tmp_path, model)
    inventory = table_source(tmp_path)
    ir.blocks = inventory.blocks()
    runner.inventory.pdf = inventory.pdf
    runner.inventory.ir = ir
    runner.inventory.mapping = inventory.mapping
    candidate, _ = await runner.read_regions(ir, {'page_idx': 0},
        RecoveryDecision(action='read_regions', reason='Read complete rows'), tmp_path / 'candidate', '')
    grid = table_structure(candidate.blocks[0])
    assert len(candidate.blocks) == 1 and grid.row_count == 4
    assert grid.source['source_rows'] == ['native:r1', 'native:r2', 'native:r3', 'native:r4']
    assert model.read_rows == [['native:r1', 'native:r2'], ['native:r1'], ['native:r2'],
                               ['native:r3', 'native:r4'], ['native:r3'], ['native:r4']]
