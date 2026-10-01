"""Sparse tables are structured evidence, not disposable OCR fragments."""
from docx import Document
from docx.oxml import OxmlElement

from app.adapters.word import _table_to_html
from app.models.document_ir import Block, BlockType
from app.pipeline.package_utils import semantic_table_to_text
from app.pipeline.stages.enrich import EnrichStage
from app.pipeline.table_structure import canonical_table_html


def render(body):
    return semantic_table_to_text(canonical_table_html(
        Block(block_id='schedule', type=BlockType.TABLE, page_idx=0, payload={'table_body': body})))


def test_sparse_native_schedule_keeps_dates_and_absolute_columns():
    rows = [
        ['期間 / 執行內容', '6月', '7月', '8月', '9月', '10月', '11月'],
        ['內部資源盤點', '6/19-6/30', '', '', '', '', ''],
        ['品牌內化培訓', '', '', '', '9/19-9/23', '', ''],
        ['結案報告會議', '', '', '', '', '', '11/20-11/30'],
    ]
    body = '<table>' + ''.join('<tr>' + ''.join(f'<td>{v}</td>' for v in row) + '</tr>' for row in rows) + '</table>'
    text = render(body)
    assert '第 2 欄：6/19-6/30' in text
    assert '第 5 欄：9/19-9/23' in text
    assert '第 7 欄：11/20-11/30' in text
    assert '低可信度' not in text
    assert '表格片段或續接資料' not in text
    assert '依原始列欄位置呈現' in text


def test_empty_months_cannot_shift_work_to_another_column():
    body = '<table><tr><td>工作</td><td>5月26日</td><td>6月5日</td><td>7月15日</td><td>8月25日</td><td>9月15日</td></tr>'
    body += '<tr><td>品牌簡報製作</td><td></td><td></td><td></td><td>初稿撰寫</td><td>完稿與優化</td></tr></table>'
    text = render(body)
    assert '第 5 欄：8月25日' in text and '第 5 欄：初稿撰寫' in text
    assert '第 6 欄：9月15日' in text and '第 6 欄：完稿與優化' in text


def test_merged_source_cell_is_emitted_once_with_its_span():
    text = render('<table><tr><td colspan="3">品牌簡介</td></tr>'
                  '<tr><td rowspan="2">同一計畫</td><td>項目一</td><td>100</td></tr>'
                  '<tr><td>項目二</td><td>100</td></tr></table>')
    assert text.count('品牌簡介') == 1 and text.count('同一計畫') == 1
    assert '第 1–3 欄：品牌簡介' in text
    assert '第 1 欄（跨 2 列）：同一計畫' in text
    assert text.count('100') == 2  # Identical values in distinct cells are real data.


def test_word_repeat_row_is_layout_metadata_not_semantic_column_header():
    doc = Document()
    t = doc.add_table(rows=3, cols=3)
    t.cell(0, 0).merge(t.cell(0, 1)).text = '114年6月19日至114年11月30日'
    t.cell(0, 2).text = '單位：元'
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for c, text in zip(t.rows[1].cells, ['政府經費', '自籌款', '合計']):
        c.text = text
    for c, text in zip(t.rows[2].cells, ['300000', '300000', '600000']):
        c.text = text
    body = _table_to_html(t)
    assert '<th' not in body
    assert 'data-repeat-header="1"' in body
    text = render(body)
    assert '第 1–2 欄：114年6月19日' in text
    assert '第 3 欄：600000' in text


def test_agent_recovery_owns_suspected_table_instead_of_duplicate_initial_reconstruction():
    body = '<table><tr><td colspan="5">' + '這是原本正確合併的公司簡介。' * 8 + '</td></tr></table>'
    b = Block(block_id='summary', type=BlockType.TABLE, page_idx=0,
              payload={'table_body': body, 'img_path': 'source.jpg'})
    assert EnrichStage(None)._get_enrichment_kind(b, None, False) is None
    assert b.payload['table_body'] == body


def test_uncertain_legacy_reconstruction_cannot_overwrite_source():
    b = Block(block_id='t', type=BlockType.TABLE, page_idx=0,
              payload={'table_body': '<table><tr><td>Source 100</td></tr></table>'})
    before = b.model_dump()
    assert not EnrichStage._apply_table_reconstruction(b, {
        'parse_was_wrong': True, 'needs_review': True, 'header': ['Guessed'], 'rows': [['Guess 100']]})
    assert b.model_dump() == before
