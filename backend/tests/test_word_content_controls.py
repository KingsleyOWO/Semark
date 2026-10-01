"""Word content controls must preserve the visible reading order and values."""
import hashlib
import json

from docx import Document
from docx.oxml import OxmlElement
from PIL import Image

from app.adapters.word import parse_word_document


def _wrap(element):
    parent = element.getparent()
    index = parent.index(element)
    control = OxmlElement("w:sdt")
    properties = OxmlElement("w:sdtPr")
    alias = OxmlElement("w:alias")
    alias.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val",
              "metadata must not become content")
    properties.append(alias)
    content = OxmlElement("w:sdtContent")
    control.extend([properties, content])
    parent.insert(index, control)
    content.append(element)
    return control


def test_content_controls_preserve_nested_blocks_runs_cells_and_images(tmp_path):
    source = tmp_path / "controlled.docx"
    picture = tmp_path / "source.png"
    Image.new("RGB", (60, 60), "blue").save(picture)
    doc = Document()
    _wrap(_wrap(doc.add_heading("年度成果", 1)._p))
    paragraph = doc.add_paragraph("營收由 ")
    _wrap(paragraph.add_run("122,878")._r)
    paragraph.add_run(" 提升至 ")
    _wrap(_wrap(paragraph.add_run("125,000")._r))
    paragraph.add_run(" 仟元")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "年度"
    table.cell(0, 1).text = "淨利率"
    table.cell(1, 0).text = "2025"
    table.cell(1, 1).text = "2.6%"
    _wrap(table.cell(1, 1).paragraphs[0]._p)
    _wrap(table.cell(1, 0)._tc)
    _wrap(table._tbl)
    image_paragraph = doc.add_paragraph()
    image_run = image_paragraph.add_run()
    image_run.add_picture(str(picture))
    _wrap(image_run._r)
    _wrap(image_paragraph._p)
    doc.add_paragraph("結案說明")
    doc.save(source)
    before = hashlib.sha256(source.read_bytes()).hexdigest()

    result = parse_word_document(source, tmp_path / "parsed")

    assert result.success
    items = json.loads(result.content_list_path.read_text())
    assert [item["type"] for item in items] == ["text", "text", "table", "image", "text"]
    assert items[0]["text"] == "年度成果"
    assert items[0]["text_level"] == 1
    assert items[1]["text"] == "營收由 122,878 提升至 125,000 仟元"
    assert "<td>2025</td>" in items[2]["table_body"]
    assert "<td>2.6%</td>" in items[2]["table_body"]
    assert (result.content_list_path.parent / items[3]["img_path"]).read_bytes() == picture.read_bytes()
    assert items[4]["text"] == "結案說明"
    assert "metadata must not become content" not in result.markdown_path.read_text()
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before
