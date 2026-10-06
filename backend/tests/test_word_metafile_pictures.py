"""Word pictures stored as metafiles (spot check 2026-10-06).

Of the 106 metafiles in the stored Word reports, 88 are EMF+ only and wrap one bitmap
(a pasted slide, map or photo), 6 are an EMF that stretches one bitmap, and 12 are
drawings (WMF, EMF text and shapes). Pillow opened none of them, so drafting sent the
model bytes it could not load; LibreOffice drew the EMF+ ones blank, so review never
saw them either. The bitmaps are decoded here, a metafile that draws anything else
is left alone, and the drawings are exported as LibreOffice draws them.
"""
import asyncio
import hashlib
import io
import json
import struct
import zipfile

from docx import Document
from PIL import Image

from app.adapters.metafile import drawable_word_copy, emfplus_bitmap, metafile_bitmap, open_picture

PARGB, ARGB = 0x000E200B, 0x0026200A
# B, G, R, A per pixel: red, green, blue, then white at half opacity (premultiplied).
PIXELS = bytes([0, 0, 255, 255, 0, 255, 0, 255, 255, 0, 0, 255, 128, 128, 128, 128])
EXPECTED = [(255, 0, 0, 255), (0, 255, 0, 255), (0, 0, 255, 255), (255, 255, 255, 128)]


def emr(kind, payload=b""):
    payload += bytes(-len(payload) % 4)
    return struct.pack("<II", kind, 8 + len(payload)) + payload


def header():
    payload = bytearray(80)
    payload[32:36] = b" EMF"  # signature at byte 40 of the record
    return emr(1, bytes(payload))


def plus(kind, flags, data=b""):
    data += bytes(-len(data) % 4)
    return struct.pack("<HHII", kind, flags, 12 + len(data), len(data)) + data


def comment(*records):
    data = b"EMF+" + b"".join(records)
    return emr(70, struct.pack("<I", len(data)) + data)


def bitmap(pixels=PIXELS, fmt=PARGB, width=2, height=2, compressed=False):
    return struct.pack("<IIiiiII", 0xDBC01002, 1, width, height, width * 4, fmt, int(compressed)) + pixels


def metafile(*objects, draws=1, extra_plus=(), extra_gdi=()):
    """Header, EMF+ comments that define the objects and draw, a GDIC comment, end of file; as Office writes it."""
    records = [plus(0x4008, 0x0800, bytes(24)), *objects]  # image attributes object 0, then the picture
    records += [plus(0x4009, 0, bytes(4))]  # clear
    records += [plus(0x401A, 0x0001, bytes(40)) for _ in range(draws)]
    records += [*extra_plus, plus(0x4002, 0)]
    multiformats = emr(70, struct.pack("<I", 32) + b"GDIC" + struct.pack("<I", 0x40000004) + bytes(24))
    return (header() + comment(plus(0x4001, 0x0001, bytes(16))) + comment(*records) + multiformats
            + b"".join(extra_gdi) + emr(14, bytes(12)))


def picture_object(data, object_id=1):
    return plus(0x4008, 0x0500 | object_id, data)


def test_an_emfplus_only_metafile_yields_the_bitmap_it_draws():
    image = emfplus_bitmap(metafile(picture_object(bitmap())))

    assert image.size == (2, 2)
    assert list(image.convert("RGBA").get_flattened_data()) == EXPECTED


def test_a_straight_alpha_bitmap_keeps_its_colours():
    pixels = bytes([0, 0, 255, 255] * 3 + [255, 255, 255, 128])
    image = emfplus_bitmap(metafile(picture_object(bitmap(pixels, fmt=ARGB))))

    assert image.convert("RGBA").getpixel((1, 1)) == (255, 255, 255, 128)


def test_an_image_object_split_over_records_is_put_back_together():
    data = bitmap()
    total = struct.pack("<I", len(data))
    chunks = [data[:20], data[20:32], data[32:]]  # record data is four-byte aligned
    records = [plus(0x4008, 0x8501, total + chunks[0]), plus(0x4008, 0x8501, total + chunks[1]),
               plus(0x4008, 0x0501, total + chunks[2])]

    image = emfplus_bitmap(metafile(*records))

    assert list(image.convert("RGBA").get_flattened_data()) == EXPECTED


def test_a_compressed_bitmap_is_the_png_or_jpeg_it_carries():
    buffer = io.BytesIO()
    Image.new("RGB", (3, 1), (10, 200, 30)).save(buffer, "PNG")

    image = emfplus_bitmap(metafile(picture_object(bitmap(buffer.getvalue(), width=3, height=1, compressed=True))))

    assert image.size == (3, 1) and image.convert("RGB").getpixel((2, 0)) == (10, 200, 30)


def test_a_metafile_that_draws_anything_besides_one_bitmap_is_left_alone():
    one = picture_object(bitmap())
    assert emfplus_bitmap(metafile(one, extra_plus=[plus(0x401C, 0, bytes(28))])) is None  # text
    assert emfplus_bitmap(metafile(one, draws=2)) is None
    assert emfplus_bitmap(metafile(one, extra_gdi=[emr(54, bytes(8))])) is None  # a GDI line
    metafile_object = struct.pack("<II", 0xDBC01002, 2) + bytes(16)
    assert emfplus_bitmap(metafile(picture_object(metafile_object))) is None


def test_unreadable_bytes_are_not_a_picture():
    whole = metafile(picture_object(bitmap()))
    assert emfplus_bitmap(b"EMF source bytes") is None
    assert emfplus_bitmap(whole[: len(whole) // 2]) is None
    assert emfplus_bitmap(metafile(picture_object(bitmap(PIXELS[:-4])))) is None


def test_open_picture_reads_ordinary_images_and_emfplus_metafiles():
    buffer = io.BytesIO()
    Image.new("RGB", (5, 4), "white").save(buffer, "PNG")
    with open_picture(buffer.getvalue()) as ordinary:
        assert ordinary.size == (5, 4)
    with open_picture(metafile(picture_object(bitmap()))) as decoded:
        assert decoded.size == (2, 2)


def _docx(tmp_path, emf=None, name="report.docx"):
    png = tmp_path / "picture.png"
    Image.new("RGB", (4, 4), "white").save(png)
    plain = tmp_path / ("plain-" + name)
    document = Document()
    document.add_paragraph("圖一 示範公司全球據點")
    document.add_picture(str(png))
    document.save(plain)
    if emf is None:
        return plain
    target = tmp_path / name
    with zipfile.ZipFile(plain) as src, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data, part = src.read(info), info.filename
            if part.startswith("word/media/") and part.endswith(".png"):
                part, data = part[:-4] + ".emf", emf
            elif part.endswith(".rels"):
                data = data.replace(b'.png"', b'.emf"')
            elif part == "[Content_Types].xml":
                data = data.replace(b"</Types>", b'<Default Extension="emf" ContentType="image/x-emf"/></Types>')
            dst.writestr(part, data)
    return target


def test_the_word_parser_exports_an_emfplus_picture_as_the_png_it_draws(tmp_path):
    from app.adapters.word import parse_word_document

    emf = metafile(picture_object(bitmap()))
    result = parse_word_document(_docx(tmp_path, emf), tmp_path / "out")

    item = next(i for i in json.loads(result.content_list_path.read_text()) if i["type"] == "image")
    assert item["img_path"].endswith(".png")
    with Image.open(result.content_list_path.parent / item["img_path"]) as exported:
        assert list(exported.convert("RGBA").get_flattened_data()) == EXPECTED
    # The source part is still identified by its own bytes, as review and recovery look it up.
    assert item["metadata"]["image_sha256"] == hashlib.sha256(emf).hexdigest()


def test_libreoffice_gets_a_copy_whose_emfplus_pictures_are_the_png_they_draw(tmp_path):
    source = _docx(tmp_path, metafile(picture_object(bitmap())))
    before = source.read_bytes()

    copy = drawable_word_copy(source, tmp_path / "drawable")

    assert copy.name == source.name and copy.parent == tmp_path / "drawable"
    assert source.read_bytes() == before
    with zipfile.ZipFile(copy) as archive:
        media = [n for n in archive.namelist() if n.startswith("word/media/")]
        assert media == ["word/media/image1.emf"]
        with Image.open(io.BytesIO(archive.read(media[0]))) as image:
            assert image.format == "PNG" and image.size == (2, 2)
    assert drawable_word_copy(_docx(tmp_path, name="ordinary.docx"), tmp_path / "other") is None


def test_word_to_pdf_conversion_reads_the_drawable_copy(tmp_path, monkeypatch):
    from app.pipeline.stages import parse

    source = _docx(tmp_path, metafile(picture_object(bitmap())))
    seen = {}

    class Done:
        returncode = 0

        async def communicate(self):
            return b"", b""

    async def libreoffice(*args, **kwargs):
        seen["input"] = args[-1]
        (tmp_path / "pdf" / "report.pdf").write_bytes(b"%PDF-1.4")
        return Done()
    monkeypatch.setattr(parse.asyncio, "create_subprocess_exec", libreoffice)

    pdf, error = asyncio.run(parse.ParseStage._convert_office_to_pdf(source, tmp_path / "pdf"))

    assert error is None and pdf == tmp_path / "pdf/report.pdf"
    assert seen["input"] == str(tmp_path / "pdf/drawable/report.docx")


def test_review_matching_and_recovery_read_an_emfplus_picture():
    from app.pipeline.recovery_evidence import _pixels
    from app.pipeline.word_mapping import _image_data

    emf = metafile(picture_object(bitmap()))

    assert _image_data(emf)["signature"][0] == (2, 2)
    assert _pixels(emf)[0] == (2, 2)


def stretch_dib(width=2, height=2, rop=0x00CC0020, source=None):
    """EMR_STRETCHDIBITS copying a bottom-up 24-bit DIB whole: blue row, then red row on top."""
    source = source or (0, 0, width, height)
    rows = (bytes([255, 0, 0] * width) + bytes(-3 * width % 4)) + (bytes([0, 0, 255] * width) + bytes(-3 * width % 4))
    info = struct.pack("<IiiHHIIiiII", 40, width, height, 1, 24, 0, len(rows), 0, 0, 0, 0)
    fixed = struct.pack("<4i6i4I4I", 0, 0, width, height, 0, 0, *source,
                        80, len(info), 80 + len(info), len(rows), 0, rop, width, height)
    return emr(81, fixed + info + rows)


def test_an_emf_that_stretches_one_bitmap_yields_that_bitmap():
    image = metafile_bitmap(header() + stretch_dib() + emr(14, bytes(12)))

    assert image.size == (2, 2)
    assert image.convert("RGB").getpixel((0, 0)) == (255, 0, 0)
    assert image.convert("RGB").getpixel((0, 1)) == (0, 0, 255)


def test_a_source_area_a_few_pixels_past_the_bitmap_still_copies_it_whole():
    # Some exporters round the source size up: 1140 x 803 for a 1139 x 803 bitmap.
    image = metafile_bitmap(header() + stretch_dib(width=200, height=2, source=(0, 0, 202, 2)) + emr(14, bytes(12)))

    assert image.size == (200, 2)
    assert metafile_bitmap(header() + stretch_dib(width=200, height=2, source=(1, 0, 199, 2)) + emr(14, bytes(12))) is None
    assert metafile_bitmap(header() + stretch_dib(width=200, height=2, source=(0, 0, 300, 2)) + emr(14, bytes(12))) is None


def test_an_emf_that_draws_more_than_one_bitmap_is_left_alone():
    assert metafile_bitmap(header() + stretch_dib() + emr(84, bytes(76)) + emr(14, bytes(12))) is None  # text
    assert metafile_bitmap(header() + stretch_dib() + stretch_dib() + emr(14, bytes(12))) is None
    assert metafile_bitmap(header() + stretch_dib(rop=0x00EE0086) + emr(14, bytes(12))) is None  # merged, not copied


def _a4_pdf_with(path, rect):
    import fitz

    with fitz.open() as document:
        page = document.new_page(width=595, height=842)
        if rect:
            page.draw_rect(fitz.Rect(*rect), color=(0, 0, 0), fill=(0, 0, 0))
        document.save(path)


def test_a_drawn_metafile_is_exported_as_libreoffice_draws_it_without_the_page_margin(tmp_path, monkeypatch):
    from app.pipeline.stages import parse

    images = tmp_path / "images"
    images.mkdir()
    (images / "docx_img0000.wmf").write_bytes(b"drawn logo")
    (images / "docx_img0001.emf").write_bytes(b"blank drawing")
    items = [{"type": "image", "img_path": "images/docx_img0000.wmf", "metadata": {"image_sha256": "abc"}},
             {"type": "image", "img_path": "images/docx_img0001.emf", "metadata": {"image_sha256": "def"}},
             {"type": "text", "text": "圖一 示範公司標誌"}]
    content_list = tmp_path / "report_content_list.json"
    content_list.write_text(json.dumps(items, ensure_ascii=False))
    markdown = tmp_path / "report.md"
    markdown.write_text("![圖一](images/docx_img0000.wmf)\n\n![圖二](images/docx_img0001.emf)\n")

    async def libreoffice(source, output_dir, **kwargs):
        output_dir.mkdir(parents=True, exist_ok=True)
        pdf = output_dir / (source.stem + ".pdf")
        _a4_pdf_with(pdf, (100, 200, 200, 250) if source.suffix == ".wmf" else None)
        return pdf, None
    monkeypatch.setattr(parse.ParseStage, "_convert_office_to_pdf", staticmethod(libreoffice))

    rendered = asyncio.run(parse.render_word_metafiles(content_list, markdown))

    after = json.loads(content_list.read_text())
    assert rendered == 1
    assert after[0]["img_path"] == "images/docx_img0000.png" and after[0]["metadata"]["image_sha256"] == "abc"
    with Image.open(tmp_path / after[0]["img_path"]) as png:
        # 100 x 50 points drawn at four pixels per point, plus the kept margin
        assert abs(png.width - (400 + 2 * parse.METAFILE_MARGIN)) <= 8
        assert abs(png.height - (200 + 2 * parse.METAFILE_MARGIN)) <= 8
    assert after[1]["img_path"] == "images/docx_img0001.emf"  # drew nothing: left as it was
    assert "images/docx_img0000.png" in markdown.read_text() and "images/docx_img0001.emf" in markdown.read_text()
    assert sorted(p.name for p in images.iterdir()) == ["docx_img0000.png", "docx_img0000.wmf", "docx_img0001.emf"]
