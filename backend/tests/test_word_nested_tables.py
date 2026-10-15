"""A table placed inside a table cell or a text box is delivered with its text.

Batch of 44 reports, 2026-10-08: a cell's text was read from its paragraphs only, and a
text box from its paragraphs only, so a table inside either was dropped whole: a patent
list and a product comparison (one report, ~350 characters), a workshop agenda (~400) and
a mystery-shopper brief in a text box (~140). The review saw the gap on the page image but
ran out of time to repair it. Each nested row is now one string, its cells joined by " | ".
"""

import json
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml

from app.adapters.word import parse_word_document

PATENTS = [["公告號", "專利名稱"], ["M500001", "工作交換台結構（一）"], ["M500002", "分度盤軸承結構"]]


def _items(tmp_path: Path, doc) -> list[dict]:
    source = tmp_path / "nested.docx"
    doc.save(source)
    result = parse_word_document(source, tmp_path / "out")
    assert result.success, result.error
    return json.loads(result.content_list_path.read_text(encoding="utf-8"))


def _fill(table, rows):
    for r, values in enumerate(rows):
        for c, value in enumerate(values):
            table.cell(r, c).text = value


def test_a_table_inside_a_cell_is_delivered_in_order_with_the_cell_text(tmp_path):
    doc = Document()
    outer = doc.add_table(rows=1, cols=2)
    outer.cell(0, 0).text = "專利布局"
    cell = outer.cell(0, 1)
    cell.paragraphs[0].text = "晨光精機專利如下："
    _fill(cell.add_table(rows=3, cols=2), PATENTS)
    cell.add_paragraph("以上均為新型專利。")

    body = next(i["table_body"] for i in _items(tmp_path, doc) if i["type"] == "table")

    assert ("晨光精機專利如下： / 公告號 | 專利名稱；M500001 | 工作交換台結構（一）；"
            "M500002 | 分度盤軸承結構 / 以上均為新型專利。") in body


def test_an_empty_corner_cell_keeps_the_columns_aligned(tmp_path):
    doc = Document()
    outer = doc.add_table(rows=1, cols=1)
    _fill(outer.cell(0, 0).add_table(rows=2, cols=3), [["", "晨光", "暮雲"], ["口味", "烏龍", "紅茶"]])

    body = next(i["table_body"] for i in _items(tmp_path, doc) if i["type"] == "table")

    assert "| 晨光 | 暮雲；口味 | 烏龍 | 紅茶" in body


def test_a_table_two_levels_down_is_delivered(tmp_path):
    doc = Document()
    outer = doc.add_table(rows=1, cols=1)
    middle = outer.cell(0, 0).add_table(rows=1, cols=2)
    middle.cell(0, 0).text = "議程"
    _fill(middle.cell(0, 1).add_table(rows=2, cols=2), [["5分", "開場與暖身"], ["40分", "永續趨勢"]])

    body = next(i["table_body"] for i in _items(tmp_path, doc) if i["type"] == "table")

    assert "5分 | 開場與暖身" in body and "40分 | 永續趨勢" in body


TEXT_BOX_TABLE_RUN = """
<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
     xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"
     xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
     xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
     xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
     xmlns:v="urn:schemas-microsoft-com:vml">
  <mc:AlternateContent>
    <mc:Choice Requires="wps">
      <w:drawing><wp:anchor><a:graphic><a:graphicData><wps:wsp><wps:txbx><w:txbxContent>
        {box}
      </w:txbxContent></wps:txbx></wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing>
    </mc:Choice>
    <mc:Fallback>
      <w:pict><v:shape><v:textbox><w:txbxContent>
        {box}
      </w:txbxContent></v:textbox></v:shape></w:pict>
    </mc:Fallback>
  </mc:AlternateContent>
</w:r>
"""
BOX = """
<w:tbl><w:tr>
  <w:tc><w:p><w:r><w:t>基本設定</w:t></w:r></w:p></w:tc>
  <w:tc><w:p><w:r><w:t>均在週五傍晚購買。</w:t></w:r></w:p><w:p><w:r><w:t>均購買兩片蛋糕。</w:t></w:r></w:p></w:tc>
</w:tr><w:tr>
  <w:tc><w:p><w:r><w:t>購買地點</w:t></w:r></w:p></w:tc>
  <w:tc><w:p><w:r><w:t>北區：晨光店</w:t></w:r></w:p></w:tc>
</w:tr></w:tbl>
<w:p><w:r><w:t>神秘客說明</w:t></w:r></w:p>
"""


def test_a_table_inside_a_text_box_is_delivered_once_per_row(tmp_path):
    doc = Document()
    doc.add_paragraph("神秘客體驗")
    host = doc.add_paragraph()
    host._p.append(parse_xml(TEXT_BOX_TABLE_RUN.format(box=BOX)))
    doc.add_paragraph("調查結果")

    texts = [i["text"] for i in _items(tmp_path, doc) if i["type"] == "text"]

    assert texts == ["神秘客體驗", "基本設定 | 均在週五傍晚購買。 / 均購買兩片蛋糕。", "購買地點 | 北區：晨光店",
                     "神秘客說明", "調查結果"]
