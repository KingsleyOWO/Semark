"""Footnotes and endnotes are delivered after the block that cites them (2026-10-16).

python-docx reads neither part. One report's seven source links reached no delivery; only
page re-reads ever found them, and those split the paragraphs around them at line wraps.
"""
import json

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import parse_xml

from app.adapters.word import parse_word_document

W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def _notes(kind, notes):
    body = "".join(f'<w:{kind} w:id="{i}"><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:{kind}>' for i, text in notes)
    separator = f'<w:{kind} w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:{kind}>'
    return f'<w:{kind}s {W}>{separator}{body}</w:{kind}s>'.encode()


def _cite(paragraph, kind, note_id):
    paragraph._p.append(parse_xml(f'<w:r {W}><w:{kind}Reference w:id="{note_id}"/></w:r>'))


def test_notes_follow_the_paragraph_that_cites_them(tmp_path):
    document = Document()
    first = document.add_paragraph("晴諾米市場規模持續成長")
    _cite(first, "footnote", 2)
    document.add_paragraph("晴川拼圖的通路策略")
    last = document.add_paragraph("結語")
    _cite(last, "endnote", 1)
    _cite(last, "footnote", 3)
    for kind, notes in (("footnote", [(2, "晴諾米研究：https://example.org/a"), (3, "https://example.org/b")]),
                        ("endnote", [(1, "晴川年報")])):
        part = Part(PackURI(f"/word/{kind}s.xml"),
                    f"application/vnd.openxmlformats-officedocument.wordprocessingml.{kind}s+xml",
                    _notes(kind, notes), document.part.package)
        document.part.relate_to(part, RT.FOOTNOTES if kind == "footnote" else RT.ENDNOTES)
    path = tmp_path / "notes.docx"
    document.save(path)

    result = parse_word_document(path, tmp_path / "out")
    texts = [item["text"] for item in json.loads(result.content_list_path.read_text()) if item["type"] == "text"]
    assert texts == ["晴諾米市場規模持續成長", "註1：晴諾米研究：https://example.org/a", "晴川拼圖的通路策略",
                     "結語", "尾註1：晴川年報", "註2：https://example.org/b"]
