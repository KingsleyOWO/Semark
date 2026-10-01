"""Text Word renders but python-docx does not hand back.

Automatic list numbers live in numbering.xml, not in the paragraph runs, and a
text box sits in w:txbxContent inside a drawing. paragraph.text returns neither,
so a source page showing 「（一） 現況評估」 or a contract number in a floating box
shipped without them.
"""

import json
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml

from app.adapters.word import parse_word_document

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _add_numbering(doc, num_id: int, levels: list[tuple[str, str]]) -> None:
    """Register an abstract numbering whose level i uses (numFmt, lvlText)."""
    numbering = doc.part.numbering_part.element
    lvls = "".join(
        f'<w:lvl w:ilvl="{i}"><w:start w:val="1"/><w:numFmt w:val="{fmt}"/>'
        f'<w:lvlText w:val="{text}"/></w:lvl>'
        for i, (fmt, text) in enumerate(levels)
    )
    abstract_id = 900 + num_id
    abstract = parse_xml(f'<w:abstractNum xmlns:w="{W_NS}" w:abstractNumId="{abstract_id}">{lvls}</w:abstractNum>')
    num = parse_xml(f'<w:num xmlns:w="{W_NS}" w:numId="{num_id}"><w:abstractNumId w:val="{abstract_id}"/></w:num>')
    first_num = numbering.find(f"{{{W_NS}}}num")
    if first_num is not None:
        first_num.addprevious(abstract)
    else:
        numbering.append(abstract)
    numbering.append(num)


def _numbered(doc, text: str, num_id: int, ilvl: int = 0) -> None:
    paragraph = doc.add_paragraph(text)
    paragraph._p.get_or_add_pPr().append(parse_xml(
        f'<w:numPr xmlns:w="{W_NS}"><w:ilvl w:val="{ilvl}"/><w:numId w:val="{num_id}"/></w:numPr>'
    ))


def _texts(tmp_path: Path, doc) -> list[str]:
    source = tmp_path / "numbered.docx"
    doc.save(source)
    result = parse_word_document(source, tmp_path / "out")
    assert result.success, result.error
    items = json.loads(result.content_list_path.read_text(encoding="utf-8"))
    return [item["text"] for item in items if item["type"] == "text"]


def test_auto_numbering_is_delivered_with_its_paragraph(tmp_path):
    doc = Document()
    _add_numbering(doc, 71, [("taiwaneseCountingThousand", "（%1）")])
    _add_numbering(doc, 72, [("decimal", "%1.")])
    _numbered(doc, "現況評估", 71)
    _numbered(doc, "內容行銷與社群媒體優化", 72)
    _numbered(doc, "數據驅動的網站優化", 72)
    _numbered(doc, "持續監測與優化", 72)

    assert _texts(tmp_path, doc) == [
        "（一） 現況評估",
        "1. 內容行銷與社群媒體優化",
        "2. 數據驅動的網站優化",
        "3. 持續監測與優化",
    ]


def test_nested_numbering_restarts_under_each_parent(tmp_path):
    doc = Document()
    _add_numbering(doc, 73, [("decimal", "%1."), ("decimal", "%1.%2.")])
    _numbered(doc, "背景", 73, 0)
    _numbered(doc, "市場", 73, 1)
    _numbered(doc, "競品", 73, 1)
    _numbered(doc, "策略", 73, 0)
    _numbered(doc, "定位", 73, 1)

    assert _texts(tmp_path, doc) == ["1. 背景", "1.1. 市場", "1.2. 競品", "2. 策略", "2.1. 定位"]


def test_numbering_inherited_from_the_paragraph_style(tmp_path):
    doc = Document()
    doc.add_paragraph("準備資料", style="List Number")
    doc.add_paragraph("送出申請", style="List Number")

    texts = _texts(tmp_path, doc)

    assert [t.split(" ", 1)[1] for t in texts] == ["準備資料", "送出申請"]
    assert texts[0].split(" ", 1)[0] != texts[1].split(" ", 1)[0]


TEXT_BOX_RUN = """
<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
     xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"
     xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
     xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
     xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
     xmlns:v="urn:schemas-microsoft-com:vml">
  <mc:AlternateContent>
    <mc:Choice Requires="wps">
      <w:drawing><wp:anchor><a:graphic><a:graphicData><wps:wsp><wps:txbx><w:txbxContent>
        <w:p><w:r><w:t>合約編號：示範001</w:t></w:r></w:p>
      </w:txbxContent></wps:txbx></wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing>
    </mc:Choice>
    <mc:Fallback>
      <w:pict><v:shape><v:textbox><w:txbxContent>
        <w:p><w:r><w:t>合約編號：示範001</w:t></w:r></w:p>
      </w:txbxContent></v:textbox></v:shape></w:pict>
    </mc:Fallback>
  </mc:AlternateContent>
</w:r>
"""


def test_text_box_content_is_delivered_once(tmp_path):
    doc = Document()
    doc.add_paragraph("結案報告書")
    host = doc.add_paragraph()
    host._p.append(parse_xml(TEXT_BOX_RUN))
    doc.add_paragraph("壹、計畫摘要")

    texts = _texts(tmp_path, doc)

    assert texts == ["結案報告書", "合約編號：示範001", "壹、計畫摘要"]


def test_text_box_inside_a_table_cell_reaches_the_cell_text(tmp_path):
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "編號"
    table.cell(0, 1).paragraphs[0]._p.append(parse_xml(TEXT_BOX_RUN))
    source = tmp_path / "cell.docx"
    doc.save(source)

    result = parse_word_document(source, tmp_path / "out")
    items = json.loads(result.content_list_path.read_text(encoding="utf-8"))
    body = next(item["table_body"] for item in items if item["type"] == "table")

    assert body.count("合約編號：示範001") == 1
