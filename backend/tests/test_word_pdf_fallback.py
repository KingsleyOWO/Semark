import io
import json
import struct
import zipfile
from pathlib import Path
from unittest.mock import AsyncMock

import pymupdf
import pytest
from docx import Document
from PIL import Image

from app.adapters.mineru import MinerUResult
from app.config import settings
from app.db.database import Database
from app.db.repositories import DocRepository
from app.models.entities import DocCreate
from app.pipeline.stages.normalize import NormalizeStage
from app.pipeline.stages.parse import ParseStage


def _word_file(path: Path, fault: str | None = None):
    document = Document()
    document.add_paragraph("Original document text")
    picture = io.BytesIO()
    Image.new("RGB", (20, 20), "blue").save(picture, format="PNG")
    picture.seek(0)
    document.add_picture(picture)
    if fault == "oversized":
        document.element.set("graphic_payload", "A" * 11_000_000)
    document.save(path)
    if fault == "crc":
        raw = io.BytesIO()
        with zipfile.ZipFile(path) as source, zipfile.ZipFile(raw, "w") as output:
            for name in source.namelist():
                output.writestr(name, source.read(name), compress_type=zipfile.ZIP_STORED)
        with zipfile.ZipFile(raw) as archive:
            info = archive.getinfo("word/media/image1.png")
        data = bytearray(raw.getvalue())
        name_len, extra_len = struct.unpack_from("<HH", data, info.header_offset + 26)
        data[info.header_offset + 30 + name_len + extra_len + 40] ^= 1
        path.write_bytes(data)


@pytest.fixture
async def word_stage(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "workspace_path", tmp_path)
    db = Database(tmp_path / "test.db")
    await db.connect()
    source = settings.get_doc_path("word-test") / "source/original.docx"
    source.parent.mkdir(parents=True)
    await DocRepository(db).create("word-test", DocCreate(
        source_path="sample.docx", ext="docx", sha256="test", size_bytes=0,
    ))
    stage = ParseStage(db)
    monkeypatch.setattr(stage.adapter, "get_version", AsyncMock(return_value="mineru-test"))
    try:
        yield stage, source
    finally:
        await db.disconnect()


def _pdf_stubs(stage, monkeypatch, *, warnings=(), mineru_error=None, empty=False):
    async def convert(source_path, output_dir, *, warnings=None):
        output_dir.mkdir(parents=True, exist_ok=True)
        pdf = output_dir / "original.pdf"
        with pymupdf.open() as document:
            document.new_page().insert_text((72, 72), "PDF fallback text")
            document.save(pdf)
        if warnings is not None:
            warnings.extend(conversion_warnings)
        return pdf, None

    async def parse(input_path, output_dir, config_override):
        assert input_path.suffix == ".pdf" and input_path.is_file()
        if mineru_error:
            return MinerUResult(success=False, output_dir=output_dir, error=mineru_error, stderr="parser diagnostic")
        parsed = output_dir / "original/auto"
        parsed.mkdir(parents=True, exist_ok=True)
        content_list = parsed / "original_content_list.json"
        items = [] if empty else [{"type": "text", "text": "PDF fallback text", "page_idx": 0}]
        content_list.write_text(json.dumps(items))
        return MinerUResult(success=True, output_dir=output_dir, content_list_path=content_list, stats={"block_count": len(items)})

    conversion_warnings = list(warnings)
    converter = AsyncMock(side_effect=convert)
    parser = AsyncMock(side_effect=parse)
    monkeypatch.setattr(stage, "_convert_office_to_pdf", converter)
    monkeypatch.setattr(stage.adapter, "parse", parser)
    return converter, parser


@pytest.mark.parametrize("fault", ["crc", "oversized"])
async def test_word_failure_uses_pdf_and_preserves_source_previews_and_cache(word_stage, monkeypatch, fault):
    stage, source = word_stage
    _word_file(source, fault)
    original = source.read_bytes()
    converter, parser = _pdf_stubs(stage, monkeypatch, warnings=["libpng error: damaged image"])

    result = await stage.run("word-test", "first")

    assert result.success, result.error
    assert result.stats["pdf_fallback"] is True
    assert result.stats["native_word_error"]
    assert result.stats["conversion_warnings"] == ["libpng error: damaged image"]
    assert source.read_bytes() == original
    assert result.content_list_path.is_file()
    origin_pdf = result.content_list_path.parent / "original_origin.pdf"
    assert origin_pdf.is_file()
    # A clean converted page takes precedence over MinerU's annotated layout.
    (origin_pdf.parent / "original_layout.pdf").write_bytes(b"annotated layout")
    normalize = NormalizeStage()
    assert normalize._find_render_pdf_path("word-test", result.content_list_path) == origin_pdf
    assert normalize._find_pdf_path("word-test", result.content_list_path) == origin_pdf
    assert not (settings.get_run_path("word-test", "first") / "temp_convert").exists()

    cached = await stage.run("word-test", "second")
    assert cached.success and cached.cache_hit
    assert cached.stats["pdf_fallback"] is True
    assert cached.stats["conversion_warnings"] == result.stats["conversion_warnings"]
    assert cached.content_list_path == result.content_list_path
    converter.assert_awaited_once()
    parser.assert_awaited_once()


async def test_readable_word_stays_native_and_uses_cache(word_stage, monkeypatch):
    stage, source = word_stage
    _word_file(source)
    converter, parser = _pdf_stubs(stage, monkeypatch)
    result = await stage.run("word-test", "first")
    cached = await stage.run("word-test", "second")
    assert result.success and result.stats["parser"] == "native_docx"
    assert cached.success and cached.cache_hit
    assert cached.content_list_path == result.content_list_path
    converter.assert_not_awaited()
    parser.assert_not_awaited()


@pytest.mark.parametrize("failure", ["conversion", "mineru", "empty"])
async def test_pdf_fallback_failure_is_not_reported_or_cached_as_success(word_stage, monkeypatch, failure):
    stage, source = word_stage
    _word_file(source, "crc")
    converter, parser = _pdf_stubs(stage, monkeypatch, mineru_error="MinerU failed" if failure == "mineru" else None, empty=failure == "empty")
    if failure == "conversion":
        converter.side_effect = None
        converter.return_value = (None, "LibreOffice unavailable")
    result = await stage.run("word-test", "first")
    assert not result.success
    assert "CRC" in result.error
    assert "PDF" in result.error
    assert not (settings.get_doc_path("word-test") / "cache/parse").exists()
    assert not (settings.get_run_path("word-test", "first") / "temp_convert").exists()
    if failure == "conversion":
        assert "LibreOffice unavailable" in result.error
        parser.assert_not_awaited()
    if failure == "mineru":
        assert "MinerU failed" in result.error


async def test_conversion_preserves_image_warnings_but_ignores_java_probe(word_stage, monkeypatch, tmp_path):
    stage, source = word_stage
    output = tmp_path / "conversion"
    output.mkdir()
    process = AsyncMock()
    process.returncode = 0

    async def communicate():
        with pymupdf.open() as document:
            document.new_page()
            document.save(output / "original.pdf")
        return b"converted", b"Warning: failed to launch javaldx - java may not function correctly\nlibpng error: bad adaptive filter value\n"

    process.communicate.side_effect = communicate
    monkeypatch.setattr("app.pipeline.stages.parse.asyncio.create_subprocess_exec", AsyncMock(return_value=process))
    warnings = []
    pdf, error = await stage._convert_office_to_pdf(source, output, warnings=warnings)
    assert error is None and pdf.is_file()
    assert warnings == ["libpng error: bad adaptive filter value"]


async def test_conversion_does_not_accept_stale_pdf(word_stage, monkeypatch, tmp_path):
    stage, source = word_stage
    output = tmp_path / "conversion"
    output.mkdir()
    (output / "original.pdf").write_bytes(b"stale output")
    process = AsyncMock()
    process.returncode = 0
    process.communicate.return_value = (b"", b"Error: source file could not be loaded")
    monkeypatch.setattr("app.pipeline.stages.parse.asyncio.create_subprocess_exec", AsyncMock(return_value=process))
    pdf, error = await stage._convert_office_to_pdf(source, output)
    assert pdf is None and error


@pytest.mark.parametrize("language", ["zh-TW", "en"])
def test_conversion_warning_stays_needs_review_after_semantic_repair(language):
    from app.pipeline.quality_gate import QualityGateResult, apply_parse_report

    gate = QualityGateResult(status="pass", score=1.0, stats={"auto_rag_ready": True})
    report = {"pdf_fallback": True, "native_word_error": "bad CRC", "conversion_warnings": ["libpng error: damaged image"]}
    apply_parse_report(gate, report, language)
    assert gate.status == "needs_review"
    assert gate.stats["auto_rag_ready"] is False
    assert gate.stats["parse"] == report
    assert gate.stats["issues_by_code"]["source_conversion_warning"] == 1
    assert gate.issues[0].evidence["warnings"] == report["conversion_warnings"]


def test_clean_pdf_fallback_does_not_require_review():
    from app.pipeline.quality_gate import QualityGateResult, apply_parse_report

    gate = QualityGateResult(status="pass", score=1.0)
    report = {"pdf_fallback": True, "native_word_error": "XML_PARSE_HUGE", "conversion_warnings": []}
    apply_parse_report(gate, report, "zh-TW")
    assert gate.status == "pass"
    assert gate.issues == []
    assert gate.stats["parse"]["pdf_fallback"] is True


@pytest.mark.parametrize("cancel", [False, True])
async def test_conversion_timeout_or_cancel_stops_process_group(word_stage, monkeypatch, tmp_path, cancel):
    import asyncio
    import signal
    from unittest.mock import Mock

    stage, source = word_stage
    process = AsyncMock()
    process.pid = 12345
    process.communicate.side_effect = [asyncio.CancelledError() if cancel else TimeoutError(), (b"", b"")]
    monkeypatch.setattr("app.pipeline.stages.parse.asyncio.create_subprocess_exec", AsyncMock(return_value=process))
    kill = Mock()
    monkeypatch.setattr("app.pipeline.stages.parse.os.killpg", kill)
    if cancel:
        with pytest.raises(asyncio.CancelledError):
            await stage._convert_office_to_pdf(source, tmp_path / "conversion")
    else:
        pdf, error = await stage._convert_office_to_pdf(source, tmp_path / "conversion")
        assert pdf is None and "timed out" in error
    kill.assert_called_once_with(12345, signal.SIGKILL)
    assert process.communicate.await_count == 2


@pytest.mark.parametrize("cancel", [False, True])
async def test_conversion_cleanup_cannot_wait_forever(word_stage, monkeypatch, tmp_path, cancel):
    import asyncio
    from unittest.mock import Mock

    stage, source = word_stage
    process = AsyncMock()
    process.pid = 12345
    process._transport = Mock()
    attempts = 0

    async def communicate():
        nonlocal attempts
        attempts += 1
        if attempts == 1:
            if cancel:
                raise asyncio.CancelledError()
            raise TimeoutError()
        await asyncio.sleep(3600)

    process.communicate.side_effect = communicate
    monkeypatch.setattr("app.pipeline.stages.parse.asyncio.create_subprocess_exec", AsyncMock(return_value=process))
    monkeypatch.setattr("app.pipeline.stages.parse.os.killpg", Mock())
    monkeypatch.setattr("app.pipeline.stages.parse.OFFICE_CLEANUP_TIMEOUT_SECONDS", 0.01)
    if cancel:
        with pytest.raises(asyncio.CancelledError):
            await asyncio.wait_for(stage._convert_office_to_pdf(source, tmp_path / "conversion"), 0.5)
    else:
        pdf, error = await asyncio.wait_for(stage._convert_office_to_pdf(source, tmp_path / "conversion"), 0.5)
        assert pdf is None and "timed out" in error
    process._transport.close.assert_called_once()

