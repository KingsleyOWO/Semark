"""Word source order and cell provenance survive parsing, normalization and mapping."""

import hashlib
import json

import fitz
from docx import Document
from PIL import Image

from app.adapters.word import parse_word_document
from app.models.document_ir import Block, BlockType
from app.pipeline.stages.enrich import EnrichStage
from app.pipeline.stages.normalize import NormalizeStage
from app.pipeline.word_mapping import locate_word_blocks, word_image_matches


def text_block(bid, text, order):
    return Block(
        block_id=bid, type=BlockType.TEXT, page_idx=0, reading_order=order, payload={"text": text}
    )


def test_repeated_word_heading_uses_neighboring_source_anchors():
    blocks = [
        text_block("a", "Unique opening", 0),
        text_block("h", "Repeated heading", 1),
        text_block("b", "Unique conclusion", 2),
    ]
    mapping, details = locate_word_blocks(
        blocks, ["repeatedheading", "uniqueopeningrepeatedheadinguniqueconclusion"]
    )
    assert mapping == {"a": [1], "h": [1], "b": [1]}
    assert details["h"]["anchor_block_ids"] == ["a", "b"]
    assert details["h"]["method"] == "ordered_source_text"


def test_order_never_invents_a_match_or_resolves_a_conflicting_interval():
    blocks = [
        text_block("a", "Later anchor", 0),
        text_block("h", "Repeated heading", 1),
        text_block("missing", "Absent words", 2),
        text_block("b", "Earlier anchor", 3),
    ]
    mapping, _ = locate_word_blocks(
        blocks, ["earlieranchorrepeatedheading", "lateranchorrepeatedheading"]
    )
    assert mapping["h"] == [] and mapping["missing"] == []
    mapping, _ = locate_word_blocks([blocks[1]], ["repeatedheading", "repeatedheading"])
    assert mapping["h"] == []


def test_order_retains_multiple_candidates_between_anchors():
    blocks = [
        text_block("a", "Start anchor", 0),
        text_block("h", "Repeated heading", 1),
        text_block("b", "End anchor", 2),
    ]
    mapping, _ = locate_word_blocks(
        blocks, ["startanchor", "repeatedheading", "repeatedheading", "endanchor"]
    )
    assert mapping["h"] == []


def test_image_caption_is_not_source_page_evidence():
    picture = Block(
        block_id="image", type=BlockType.IMAGE, page_idx=0, payload={"caption": "Unique caption"}
    )
    mapping, _ = locate_word_blocks([picture], ["uniquecaption"])
    assert mapping["image"] == []


def test_repeated_picture_uses_parent_table_pages_without_claiming_certification():
    table = Block(
        block_id="table",
        type=BlockType.TABLE,
        page_idx=0,
        reading_order=0,
        payload={
            "table_body": "<table><tr><td>Distinctive amount</td><td>Revenue 100</td></tr></table>"
        },
    )
    picture = Block(
        block_id="image",
        type=BlockType.IMAGE,
        page_idx=0,
        reading_order=1,
        payload={"word_source": {"parent_table_block_id": "table"}},
    )
    mapping, details = locate_word_blocks(
        [table, picture],
        ["distinctiveamountrevenue100", "other"],
        image_matches={"image": [{"page_idx": 0}, {"page_idx": 1}]},
    )
    assert mapping["image"] == [0]
    assert details["image"]["method"] == "ordered_image_pixels"
    assert "verified" not in details["image"]


def test_table_pictures_keep_occurrences_and_merged_cell_grid_position(tmp_path):
    picture = tmp_path / "picture.png"
    Image.new("RGB", (40, 40), "green").save(picture)
    doc = Document()
    table = doc.add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Brand information"
    table.cell(0, 2).paragraphs[0].add_run().add_picture(str(picture))
    cell = table.cell(1, 1).merge(table.cell(1, 2))
    table.cell(1, 0).text = "Brand photographs"
    for _ in range(2):
        cell.paragraphs[0].add_run().add_picture(str(picture))
    source = tmp_path / "table.docx"
    doc.save(source)
    before = source.read_bytes()
    parsed = parse_word_document(source, tmp_path / "parse")
    assert parsed.success
    items = json.loads(parsed.content_list_path.read_text())
    assert [i["type"] for i in items] == ["table", "image", "image", "image"]
    assert parsed.stats["image_count"] == 3
    assert len({i["img_path"] for i in items[1:]}) == 1
    assert [i["metadata"]["cell_path"][0]["column"] for i in items[1:]] == [2, 1, 1]
    assert items[2]["metadata"]["row_label"] == "Brand photographs"
    assert 'colspan="2"' in items[0]["table_body"]
    assert source.read_bytes() == before
    blocks = [NormalizeStage()._parse_block(item, index) for index, item in enumerate(items)]
    assert blocks[2].payload["word_source"]["parent_table_block_id"] == blocks[0].block_id
    assert EnrichStage(None)._get_enrichment_kind(blocks[2], None, False) == "figure_description"


def test_picture_only_nested_table_is_not_dropped(tmp_path):
    picture = tmp_path / "picture.png"
    Image.new("RGB", (40, 40), "blue").save(picture)
    doc = Document()
    outer = doc.add_table(rows=1, cols=1)
    inner = outer.cell(0, 0).add_table(rows=1, cols=2)
    inner.cell(0, 1).paragraphs[0].add_run().add_picture(str(picture))
    source = tmp_path / "nested.docx"
    doc.save(source)
    parsed = parse_word_document(source, tmp_path / "parse")
    assert parsed.success
    items = json.loads(parsed.content_list_path.read_text())
    assert [i["type"] for i in items] == ["table", "image"]
    path = items[1]["metadata"]["cell_path"]
    assert [(c["row"], c["column"]) for c in path] == [(0, 0), (0, 1)]
    assert items[0]["metadata"]["row_count"] == 1
    assert items[0]["metadata"]["column_count"] == 1


def test_image_mapping_requires_original_part_hash_and_painted_pdf_image(tmp_path):
    picture = tmp_path / "picture.png"
    Image.new("RGB", (40, 40), "green").save(picture)
    doc = Document()
    doc.add_picture(str(picture))
    source = tmp_path / "source.docx"
    doc.save(source)
    parsed = parse_word_document(source, tmp_path / "parse")
    items = json.loads(parsed.content_list_path.read_text())
    blocks = [NormalizeStage()._parse_block(item, index) for index, item in enumerate(items)]
    pdfpath = tmp_path / "source.pdf"
    with fitz.open() as pdf:
        pdf.new_page()
        page = pdf.new_page()
        page.insert_image(fitz.Rect(20, 20, 60, 60), filename=str(picture))
        pdf.save(pdfpath)
    matches = word_image_matches(blocks, source, pdfpath)
    assert matches[blocks[0].block_id][0]["page_idx"] == 1
    assert (
        blocks[0].payload["word_source"]["image_sha256"]
        == hashlib.sha256(picture.read_bytes()).hexdigest()
    )
    blocks[0].payload["word_source"]["image_sha256"] = "wrong hash"
    assert word_image_matches(blocks, source, pdfpath) == {}


def _pattern_picture(path):
    from PIL import ImageDraw

    image = Image.new("RGB", (240, 120), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((10, 10, 110, 100), fill="blue")
    draw.ellipse((120, 20, 220, 115), fill="orange")
    image.save(path)
    return image


def test_recompressed_picture_maps_but_close_multiple_pages_stay_ambiguous(tmp_path):
    from PIL import ImageEnhance

    picture = tmp_path / "picture.png"
    original = _pattern_picture(picture)
    doc = Document()
    doc.add_picture(str(picture))
    source = tmp_path / "source.docx"
    doc.save(source)
    parsed = parse_word_document(source, tmp_path / "parse")
    items = json.loads(parsed.content_list_path.read_text())
    blocks = [NormalizeStage()._parse_block(item, index) for index, item in enumerate(items)]
    compressed = tmp_path / "compressed.jpg"
    original.resize((216, 108), Image.Resampling.LANCZOS).save(compressed, quality=95)
    pdfpath = tmp_path / "source.pdf"
    with fitz.open() as pdf:
        page = pdf.new_page()
        page.insert_image(fitz.Rect(20, 20, 260, 140), filename=str(compressed))
        pdf.save(pdfpath)
    matches = word_image_matches(blocks, source, pdfpath)
    assert matches[blocks[0].block_id][0]["method"] == "resampled_pixels"
    mapping, _ = locate_word_blocks(blocks, [""], image_matches=matches)
    assert mapping[blocks[0].block_id] == [0]

    other = tmp_path / "other.jpg"
    ImageEnhance.Brightness(original.resize((216, 108))).enhance(0.99).save(other, quality=95)
    with fitz.open(pdfpath) as pdf:
        page = pdf.new_page()
        page.insert_image(fitz.Rect(20, 20, 260, 140), filename=str(other))
        pdf.save(tmp_path / "ambiguous.pdf")
    matches = word_image_matches(blocks, source, tmp_path / "ambiguous.pdf")
    mapping, _ = locate_word_blocks(blocks, ["", ""], image_matches=matches)
    assert mapping[blocks[0].block_id] == []


def test_image_transparency_cannot_false_match_identical_rgb_bytes(tmp_path):
    picture = tmp_path / "transparent.png"
    Image.new("RGBA", (40, 40), (0, 128, 0, 0)).save(picture)
    opaque = tmp_path / "opaque.png"
    Image.new("RGB", (40, 40), (0, 128, 0)).save(opaque)
    doc = Document()
    doc.add_picture(str(picture))
    source = tmp_path / "source.docx"
    doc.save(source)
    parsed = parse_word_document(source, tmp_path / "parse")
    items = json.loads(parsed.content_list_path.read_text())
    blocks = [NormalizeStage()._parse_block(item, index) for index, item in enumerate(items)]
    with fitz.open() as pdf:
        page = pdf.new_page()
        page.insert_image(fitz.Rect(20, 20, 60, 60), filename=str(opaque))
        pdf.save(tmp_path / "source.pdf")
    assert word_image_matches(blocks, source, tmp_path / "source.pdf") == {}


def test_floating_images_follow_visual_rows_only_in_comparable_frames(tmp_path):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    from app.adapters.word import _paragraph_image_rel_ids

    picture = tmp_path / "picture.png"
    Image.new("RGB", (40, 40), "blue").save(picture)
    doc = Document()
    paragraph = doc.add_paragraph()
    anchors = []
    for rel, x, y in [
        ("right", 300, 90),
        ("left", 10, 100),
        ("middle", 150, 95),
        ("next", 10, 300),
    ]:
        inline = paragraph.add_run().add_picture(str(picture))._inline
        inline.tag = qn("wp:anchor")
        inline.find(qn("wp:extent")).set("cy", "100")
        for tag, frame, offset in [("positionH", "column", x), ("positionV", "paragraph", y)]:
            position = OxmlElement("wp:" + tag)
            position.set("relativeFrom", frame)
            value = OxmlElement("wp:posOffset")
            value.text = str(offset)
            position.append(value)
            inline.insert(0, position)
        next(inline.iter(qn("a:blip"))).set(qn("r:embed"), rel)
        anchors.append(inline)
    assert _paragraph_image_rel_ids(paragraph) == ["left", "middle", "right", "next"]
    anchors[0].find(qn("wp:positionH")).set("relativeFrom", "page")
    assert _paragraph_image_rel_ids(paragraph) == ["right", "left", "middle", "next"]
    anchors[0].tag = qn("wp:inline")
    assert _paragraph_image_rel_ids(paragraph) == ["right", "left", "middle", "next"]


def test_word_original_image_evidence_is_pixel_faithful_and_hash_bound(tmp_path):
    picture = tmp_path / "picture.png"
    original = _pattern_picture(picture)
    doc = Document()
    doc.add_picture(str(picture))
    source = tmp_path / "source.docx"
    doc.save(source)
    parsed = parse_word_document(source, tmp_path / "parse")
    items = json.loads(parsed.content_list_path.read_text())
    blocks = [NormalizeStage()._parse_block(item, index) for index, item in enumerate(items)]
    pdfpath = tmp_path / "source.pdf"
    with fitz.open() as pdf:
        page = pdf.new_page()
        page.insert_image(fitz.Rect(20, 20, 260, 140), filename=str(picture))
        pdf.save(pdfpath)
    evidence = tmp_path / "evidence"
    evidence.mkdir()
    matches = word_image_matches(blocks, source, pdfpath, evidence_dir=evidence)
    path = evidence / matches[blocks[0].block_id][0]["source_image"]
    with Image.open(path) as exported:
        assert exported.convert("RGB").tobytes() == original.tobytes()
    blocks[0].payload["word_source"]["image_sha256"] = "incorrect"
    assert word_image_matches(blocks, source, pdfpath, evidence_dir=evidence) == {}
