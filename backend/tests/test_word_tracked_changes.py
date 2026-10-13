"""Word tracked changes are read as accepted: inserted text kept, deleted text and rows gone."""
import json
import zipfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.adapters.metafile import drawable_word_copy
from app.adapters.word import parse_word_document


def _revision(tag):
    element = OxmlElement(tag)
    element.set(qn("w:id"), "1")
    element.set(qn("w:author"), "Reviewer")
    return element


def _inserted(paragraph, text):
    """Append ``text`` to ``paragraph`` as typed with change tracking on."""
    insertion = _revision("w:ins")
    insertion.append(paragraph.add_run(text)._r)
    paragraph._p.append(insertion)


def _deleted(paragraph, text):
    deletion = _revision("w:del")
    run = OxmlElement("w:r")
    deleted = OxmlElement("w:delText")
    deleted.text = text
    run.append(deleted)
    deletion.append(run)
    paragraph._p.append(deletion)


def _tracked_docx(path):
    doc = Document()
    total = doc.add_paragraph("合計 ")
    _inserted(total, "3")
    _deleted(total, "6")
    total.add_run("00,000 元")
    _inserted(doc.add_paragraph(), "訪談對象包含四位海外業務")
    _deleted(doc.add_paragraph(), "舊版說明不應出現")
    moved = doc.add_paragraph("移動後：")
    move_to = _revision("w:moveTo")
    move_to.append(moved.add_run("第二季")._r)
    moved._p.append(move_to)
    move_from = _revision("w:moveFrom")
    move_from.append(doc.add_paragraph("第一季")._p.r_lst[0])
    doc.paragraphs[-1]._p.append(move_from)

    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "項目"
    table.cell(0, 1).text = "金額"
    table.cell(1, 0).text = "業務費"
    cell = table.cell(1, 1).paragraphs[0]
    _inserted(cell, "600,000")
    _deleted(cell, "100")
    table.cell(2, 0).text = "刪除的列"
    table.cell(2, 1).text = "999"
    row_properties = table.rows[2]._tr.get_or_add_trPr()
    row_properties.append(_revision("w:del"))
    doc.save(path)
    return path


def test_the_parser_reads_tracked_changes_as_accepted(tmp_path):
    source = _tracked_docx(tmp_path / "tracked.docx")
    before = source.read_bytes()

    result = parse_word_document(source, tmp_path / "parsed")

    assert result.success
    assert result.stats["tracked_changes_accepted"] == 9
    items = json.loads(result.content_list_path.read_text())
    texts = [item["text"] for item in items if item["type"] == "text"]
    assert texts == ["合計 300,000 元", "訪談對象包含四位海外業務", "移動後：第二季"]
    table = next(item["table_body"] for item in items if item["type"] == "table")
    assert "<td>600,000</td>" in table
    assert "100" not in table and "刪除的列" not in table and "999" not in table
    markdown = result.markdown_path.read_text()
    assert "舊版說明不應出現" not in markdown and "第一季" not in markdown
    assert source.read_bytes() == before


def test_libreoffice_draws_the_accepted_text(tmp_path):
    source = _tracked_docx(tmp_path / "tracked.docx")
    before = source.read_bytes()

    copy = drawable_word_copy(source, tmp_path / "drawable")

    assert copy is not None and source.read_bytes() == before
    with zipfile.ZipFile(copy) as archive:
        xml = archive.read("word/document.xml").decode()
    assert "w:ins " not in xml and "w:del " not in xml and "w:delText" not in xml
    assert "600,000" in xml and "訪談對象包含四位海外業務" in xml and "第二季" in xml
    assert "舊版說明不應出現" not in xml and "999" not in xml and "第一季" not in xml
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(copy) as converted:
        assert original.namelist() == converted.namelist()


def test_a_document_without_tracked_changes_is_read_as_before(tmp_path):
    source = tmp_path / "plain.docx"
    doc = Document()
    doc.add_paragraph("營收由 122,878 提升至 125,000 仟元")
    doc.save(source)

    result = parse_word_document(source, tmp_path / "parsed")

    assert result.stats["tracked_changes_accepted"] == 0
    assert drawable_word_copy(source, tmp_path / "drawable") is None
